"""
snaglist_pro.batch_processor -- Batch ZIP Processing
=====================================================
Process multiple WhatsApp chat export ZIP files through the snaglist pipeline.
"""

import os
import re
import logging
from typing import List, Dict, Optional

from snaglist_pro.pipeline import SnaglistPipeline

logger = logging.getLogger(__name__)


def _slugify(name: str) -> str:
    text = name.lower().strip()
    text = re.sub(r"[^\w\s-]", "", text)
    text = re.sub(r"[\s_]+", "-", text)
    return text.strip("-")


def _project_name_from_zip(zip_path: str) -> str:
    base = os.path.splitext(os.path.basename(zip_path))[0]
    cleaned = re.sub(r"[-_.]chat", "", base, flags=re.IGNORECASE)
    cleaned = re.sub(r"[-_.]export", "", cleaned, flags=re.IGNORECASE)
    parts = re.split(r"[-_]", cleaned)
    meaningful = [p for p in parts if len(p) > 2]
    return " ".join(meaningful).strip() if meaningful else cleaned


def _generate_master_excel(results: List[dict], output_path: str) -> None:
    try:
        from openpyxl import Workbook
        from openpyxl.styles import Font, PatternFill, Border, Side, Alignment
    except ImportError:
        logger.error("openpyxl is required for master summary Excel")
        return

    wb = Workbook()

    # ---- Overview Sheet ----
    ws_overview = wb.active
    ws_overview.title = "Overview"
    ws_overview.merge_cells("A1:F1")
    title_cell = ws_overview["A1"]
    title_cell.value = "Master Batch Summary"
    title_cell.font = Font(name="Calibri", bold=True, size=14, color="2F5496")
    title_cell.alignment = Alignment(horizontal="center")

    headers = ["Project", "Total Snags", "Matched", "Unmatched", "Open", "Closed"]
    header_fill = PatternFill(start_color="2F5496", end_color="2F5496", fill_type="solid")
    header_font = Font(name="Calibri", bold=True, color="FFFFFF", size=10)
    thin_border = Border(
        left=Side(style="thin"), right=Side(style="thin"),
        top=Side(style="thin"), bottom=Side(style="thin"),
    )

    for col_idx, h in enumerate(headers, 1):
        cell = ws_overview.cell(row=2, column=col_idx, value=h)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal="center")
        cell.border = thin_border

    grand_total = 0
    grand_matched = 0
    grand_unmatched = 0
    grand_open = 0
    grand_closed = 0

    for i, res in enumerate(results):
        row = 3 + i
        stats = res.get("stats", {})
        project = res.get("project_name", os.path.basename(res.get("zip_path", "")))
        total = stats.get("total", 0)
        matched = stats.get("matched", 0)
        unmatched = stats.get("unmatched", 0)
        open_count = stats.get("open", 0)
        closed_count = stats.get("closed", 0)

        grand_total += total
        grand_matched += matched
        grand_unmatched += unmatched
        grand_open += open_count
        grand_closed += closed_count

        vals = [project, total, matched, unmatched, open_count, closed_count]
        for col_idx, v in enumerate(vals, 1):
            cell = ws_overview.cell(row=row, column=col_idx, value=v)
            cell.font = Font(name="Calibri", size=10)
            cell.alignment = Alignment(horizontal="center")
            cell.border = thin_border

    # Grand total row
    total_row = 3 + len(results)
    grand_vals = ["TOTAL", grand_total, grand_matched, grand_unmatched, grand_open, grand_closed]
    for col_idx, v in enumerate(grand_vals, 1):
        cell = ws_overview.cell(row=total_row, column=col_idx, value=v)
        cell.font = Font(name="Calibri", bold=True, size=10)
        cell.alignment = Alignment(horizontal="center")
        cell.border = thin_border

    ws_overview.column_dimensions["A"].width = 28
    ws_overview.column_dimensions["B"].width = 14
    ws_overview.column_dimensions["C"].width = 14
    ws_overview.column_dimensions["D"].width = 14
    ws_overview.column_dimensions["E"].width = 14
    ws_overview.column_dimensions["F"].width = 14

    # ---- Quick Stats Sheet ----
    ws_stats = wb.create_sheet("Quick Stats")

    ws_stats.merge_cells("A1:C1")
    ws_stats["A1"].value = "Quick Stats"
    ws_stats["A1"].font = Font(name="Calibri", bold=True, size=14, color="2F5496")

    categories: Dict[str, int] = {}
    priorities: Dict[str, int] = {}
    vendors: Dict[str, int] = {}

    for res in results:
        stats = res.get("stats", {})
        for cat, cnt in stats.get("categories", {}).items():
            categories[cat] = categories.get(cat, 0) + cnt
        for prio, cnt in stats.get("priorities", {}).items():
            priorities[prio] = priorities.get(prio, 0) + cnt
        for vend, cnt in stats.get("vendors", {}).items():
            vendors[vend] = vendors.get(vend, 0) + cnt

    row = 3
    for title, data in [
        ("Categories", categories),
        ("Priorities", priorities),
        ("Vendors", vendors),
    ]:
        ws_stats.cell(row=row, column=1, value=title)
        ws_stats.cell(row=row, column=1).font = Font(name="Calibri", bold=True, size=11, color="2F5496")
        row += 1
        ws_stats.cell(row=row, column=1, value="Name")
        ws_stats.cell(row=row, column=2, value="Count")
        for col in (1, 2):
            ws_stats.cell(row=row, column=col).font = header_font
            ws_stats.cell(row=row, column=col).fill = header_fill
            ws_stats.cell(row=row, column=col).alignment = Alignment(horizontal="center")
            ws_stats.cell(row=row, column=col).border = thin_border
        row += 1
        for name, cnt in sorted(data.items(), key=lambda x: x[1], reverse=True):
            ws_stats.cell(row=row, column=1, value=name).border = thin_border
            ws_stats.cell(row=row, column=2, value=cnt).border = thin_border
            row += 1
        row += 1

    ws_stats.column_dimensions["A"].width = 22
    ws_stats.column_dimensions["B"].width = 14

    wb.save(output_path)
    logger.info("Master summary saved to %s", output_path)


def process_batch(
    zip_paths: List[str],
    checklist_path: Optional[str] = None,
    output_base: Optional[str] = None,
) -> List[dict]:
    if not zip_paths:
        logger.warning("No ZIP paths provided to process_batch")
        return []

    if output_base:
        os.makedirs(output_base, exist_ok=True)
    else:
        output_base = os.getcwd()

    results: List[dict] = []

    for zip_path in zip_paths:
        if not os.path.isfile(zip_path):
            logger.error("ZIP file not found: %s", zip_path)
            results.append({
                "zip_path": zip_path,
                "success": False,
                "error": f"File not found: {zip_path}",
            })
            continue

        project_name = _project_name_from_zip(zip_path)
        slug = _slugify(project_name)
        project_dir = os.path.join(output_base, slug)
        os.makedirs(project_dir, exist_ok=True)

        output_path = os.path.join(project_dir, f"{slug}_snaglist.xlsx")

        logger.info("Processing %s into %s ...", zip_path, project_dir)

        try:
            pipeline = SnaglistPipeline()
            pipeline.run(
                zip_path=zip_path,
                checklist_path=checklist_path or "",
                output_dir=os.path.dirname(output_path),
            )

            stats = getattr(pipeline, "stats", {})
            if not stats:
                stats = _compute_stats_from_pipeline(pipeline)

            results.append({
                "zip_path": zip_path,
                "project_name": project_name,
                "slug": slug,
                "output_dir": project_dir,
                "output_path": output_path,
                "success": True,
                "stats": stats,
            })
            logger.info("Successfully processed %s -> %s", zip_path, output_path)

        except Exception as exc:
            logger.exception("Failed to process %s: %s", zip_path, exc)
            results.append({
                "zip_path": zip_path,
                "project_name": project_name,
                "slug": slug,
                "output_dir": project_dir,
                "success": False,
                "error": str(exc),
            })

    # Generate master summary
    master_path = os.path.join(output_base, "_master_summary.xlsx")
    _generate_master_excel(results, master_path)

    return results


def _compute_stats_from_pipeline(pipeline) -> dict:
    stats = {
        "total": 0,
        "matched": 0,
        "unmatched": 0,
        "open": 0,
        "closed": 0,
        "categories": {},
        "priorities": {},
        "vendors": {},
    }

    snags = getattr(pipeline, "snags", []) or []
    stats["total"] = len(snags)

    for snag in snags:
        if isinstance(snag, dict):
            if snag.get("match_score", 0) > 0:
                stats["matched"] += 1
            else:
                stats["unmatched"] += 1

            status = snag.get("status", "Open")
            if status.lower() in ("closed", "resolved"):
                stats["closed"] += 1
            else:
                stats["open"] += 1

            cat = snag.get("category", "Unassigned")
            stats["categories"][cat] = stats["categories"].get(cat, 0) + 1

            prio = snag.get("priority", "Medium")
            stats["priorities"][prio] = stats["priorities"].get(prio, 0) + 1

            vend = snag.get("vendor", "")
            if vend:
                stats["vendors"][vend] = stats["vendors"].get(vend, 0) + 1

    return stats
