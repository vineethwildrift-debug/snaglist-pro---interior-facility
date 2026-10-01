import logging
from typing import Dict

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Border, Side, Alignment, numbers
from openpyxl.chart import BarChart, PieChart, Reference
from openpyxl.formatting.rule import DataBarRule
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.dimensions import ColumnDimension

logger = logging.getLogger(__name__)

_HEADER_FILL = PatternFill(start_color="2F5496", end_color="2F5496", fill_type="solid")
_HEADER_FONT = Font(name="Calibri", bold=True, color="FFFFFF", size=11)
_TITLE_FONT = Font(name="Calibri", bold=True, size=14, color="2F5496")
_SECTION_FONT = Font(name="Calibri", bold=True, size=12, color="2F5496")
_BODY_FONT = Font(name="Calibri", size=10)
_BOLD_FONT = Font(name="Calibri", bold=True, size=10)
_THIN_BORDER = Border(
    left=Side(style="thin"), right=Side(style="thin"),
    top=Side(style="thin"), bottom=Side(style="thin"),
)
_CENTER = Alignment(horizontal="center", vertical="center")
_LEFT = Alignment(horizontal="left", vertical="center")


def _style_header(ws, row, col_start, col_end):
    for col in range(col_start, col_end + 1):
        c = ws.cell(row=row, column=col)
        c.font = _HEADER_FONT
        c.fill = _HEADER_FILL
        c.alignment = _CENTER
        c.border = _THIN_BORDER


def _style_body(ws, row, col_start, col_end):
    for col in range(col_start, col_end + 1):
        c = ws.cell(row=row, column=col)
        c.font = _BODY_FONT
        c.border = _THIN_BORDER
        c.alignment = _CENTER


def _style_bold(ws, row, col_start, col_end):
    for col in range(col_start, col_end + 1):
        c = ws.cell(row=row, column=col)
        c.font = _BOLD_FONT
        c.border = _THIN_BORDER
        c.alignment = _CENTER


def add_summary_sheet(wb: Workbook, stats: dict, data_sheet: str = "") -> None:
    if "Summary" in wb.sheetnames:
        del wb["Summary"]

    ws = wb.create_sheet("Summary")
    ws.sheet_properties.tabColor = "2F5496"

    ref = f"'{data_sheet}'" if data_sheet else "Sheet1"

    total = stats.get("total", 0)
    total_rows = stats.get("total_rows", total)
    matched = stats.get("matched", 0)
    unmatched = stats.get("unmatched", 0)
    open_count = stats.get("open", 0)
    closed_count = stats.get("closed", 0)
    categories: Dict = stats.get("categories", {})
    priorities: Dict = stats.get("priorities", {})
    vendors: Dict = stats.get("vendors", {})

    last_data_row = 7 + total_rows
    status_range = f"{ref}!I8:I{last_data_row}" if data_sheet else f"I8:I{last_data_row}"
    # Column J holds the Snag Points text. Restricting the counts to rows that
    # have one keeps "Open"/"Closed" describing real snags instead of also
    # counting checklist questions that had nothing reported against them.
    snag_range = f"{ref}!J8:J{last_data_row}" if data_sheet else f"J8:J{last_data_row}"

    ws.merge_cells("A1:H1")
    ws["A1"].value = "Snaglist Summary Dashboard"
    ws["A1"].font = _TITLE_FONT
    ws["A1"].alignment = Alignment(horizontal="center")

    ws.merge_cells("A3:B3")
    ws["A3"].value = "Key Metrics"
    ws["A3"].font = _SECTION_FONT

    ws.cell(row=4, column=1, value="Metric").font = _HEADER_FONT
    ws.cell(row=4, column=1).fill = _HEADER_FILL
    ws.cell(row=4, column=1).alignment = _CENTER
    ws.cell(row=4, column=1).border = _THIN_BORDER
    ws.cell(row=4, column=2, value="Value").font = _HEADER_FONT
    ws.cell(row=4, column=2).fill = _HEADER_FILL
    ws.cell(row=4, column=2).alignment = _CENTER
    ws.cell(row=4, column=2).border = _THIN_BORDER

    unmatched_snags = stats.get("unmatched_snags", 0)
    checklist_items = stats.get("checklist_items", 0)
    unmatched_checklist = stats.get("unmatched_checklist", max(unmatched - unmatched_snags, 0))

    # Row map (the Match Rate formula and the category percentages key off
    # these cells):
    #   B5 Total Snags   B6 Matched   B7 Unmatched Snags   B8 Checklist Items
    #   B9 Match Rate    B10 Open     B11 Closed
    metrics_data = [
        ("Total Snags", total, False),
        ("Matched to Checklist", matched, False),
        ("Unmatched Snags", unmatched_snags, False),
        ("Checklist Items", checklist_items or total_rows, False),
        # Match rate is matched / (matched + unmatched snags) - the share of
        # extracted snags that the checklist accounted for.
        ("Match Rate", "=B6/(B6+B7)" if (matched + unmatched_snags) > 0 else 0, True),
        ("Open", open_count, False),
        ("Closed", closed_count, False),
        ("Checklist Items With No Snag", unmatched_checklist, False),
        ("Report Rows Printed", total_rows, False),
    ]
    for i, (label, val, is_formula) in enumerate(metrics_data):
        r = 5 + i
        ws.cell(row=r, column=1, value=label).font = _BODY_FONT
        ws.cell(row=r, column=1).border = _THIN_BORDER
        ws.cell(row=r, column=1).alignment = _LEFT
        c = ws.cell(row=r, column=2)
        if is_formula:
            c.value = val
            c.number_format = "0%"
        else:
            c.value = val
        c.font = _BOLD_FONT
        c.border = _THIN_BORDER
        c.alignment = _CENTER

    ws.merge_cells("D3:E3")
    ws["D3"].value = "Status Overview"
    ws["D3"].font = _SECTION_FONT

    ws.cell(row=4, column=4, value="Status").font = _HEADER_FONT
    ws.cell(row=4, column=4).fill = _HEADER_FILL
    ws.cell(row=4, column=4).alignment = _CENTER
    ws.cell(row=4, column=4).border = _THIN_BORDER
    ws.cell(row=4, column=5, value="Count").font = _HEADER_FONT
    ws.cell(row=4, column=5).fill = _HEADER_FILL
    ws.cell(row=4, column=5).alignment = _CENTER
    ws.cell(row=4, column=5).border = _THIN_BORDER

    status_formulas = [
        ("Open", f'=COUNTIFS({status_range},"Open",{snag_range},"<>")'),
        ("Closed", f'=COUNTIFS({status_range},"Closed",{snag_range},"<>")'),
    ]
    for i, (label, formula) in enumerate(status_formulas):
        r = 5 + i
        ws.cell(row=r, column=4, value=label).font = _BODY_FONT
        ws.cell(row=r, column=4).border = _THIN_BORDER
        ws.cell(row=r, column=5).value = formula
        ws.cell(row=r, column=5).font = _BOLD_FONT
        ws.cell(row=r, column=5).border = _THIN_BORDER
        ws.cell(row=r, column=5).alignment = _CENTER

    cat_start = 16  # clears the nine Key Metrics rows (5-13)
    ws.merge_cells(f"A{cat_start}:E{cat_start}")
    ws.cell(row=cat_start, column=1, value="Category Distribution").font = _SECTION_FONT

    ws.cell(row=cat_start + 1, column=1, value="Category").font = _HEADER_FONT
    ws.cell(row=cat_start + 1, column=1).fill = _HEADER_FILL
    ws.cell(row=cat_start + 1, column=1).alignment = _CENTER
    ws.cell(row=cat_start + 1, column=1).border = _THIN_BORDER
    ws.cell(row=cat_start + 1, column=2, value="Count").font = _HEADER_FONT
    ws.cell(row=cat_start + 1, column=2).fill = _HEADER_FILL
    ws.cell(row=cat_start + 1, column=2).alignment = _CENTER
    ws.cell(row=cat_start + 1, column=2).border = _THIN_BORDER
    ws.cell(row=cat_start + 1, column=3, value="%").font = _HEADER_FONT
    ws.cell(row=cat_start + 1, column=3).fill = _HEADER_FILL
    ws.cell(row=cat_start + 1, column=3).alignment = _CENTER
    ws.cell(row=cat_start + 1, column=3).border = _THIN_BORDER

    cat_row = cat_start + 2
    sorted_cats = sorted(categories.items(), key=lambda x: x[1], reverse=True)
    for i, (name, count) in enumerate(sorted_cats):
        r = cat_row + i
        ws.cell(row=r, column=1, value=name)
        ws.cell(row=r, column=2, value=count)
        ws.cell(row=r, column=3, value=f"=B{r}/$B$5")
        ws.cell(row=r, column=3).number_format = "0%"
        _style_body(ws, r, 1, 3)

    if not categories:
        ws.cell(row=cat_row, column=1, value="No categories").font = _BODY_FONT
        ws.cell(row=cat_row, column=1).border = _THIN_BORDER

    if categories:
        chart_bar = BarChart()
        chart_bar.type = "col"
        chart_bar.title = "Snags by Category"
        chart_bar.y_axis.title = "Count"
        chart_bar.x_axis.title = "Category"
        chart_bar.style = 10
        data_ref = Reference(ws, min_col=2, min_row=cat_start + 1, max_row=cat_row + len(categories) - 1)
        cats_ref = Reference(ws, min_col=1, min_row=cat_row, max_row=cat_row + len(categories) - 1)
        chart_bar.add_data(data_ref, titles_from_data=True)
        chart_bar.set_categories(cats_ref)
        chart_bar.shape = 4
        ws.add_chart(chart_bar, f"E{cat_start}")

        last_cat_row = cat_row + len(categories) - 1
        ws.cell(row=last_cat_row + 1, column=1, value="TOTAL").font = _BOLD_FONT
        ws.cell(row=last_cat_row + 1, column=1).border = _THIN_BORDER
        ws.cell(row=last_cat_row + 1, column=2, value=f"=SUM(B{cat_row}:B{last_cat_row})")
        ws.cell(row=last_cat_row + 1, column=2).font = _BOLD_FONT
        ws.cell(row=last_cat_row + 1, column=2).border = _THIN_BORDER
        ws.cell(row=last_cat_row + 1, column=2).alignment = _CENTER

        if len(sorted_cats) > 1:
            end = last_cat_row - cat_row
            ws.conditional_formatting.add(
                f"B{cat_row}:B{last_cat_row}",
                DataBarRule(start_type="min", end_type="max", color="5B9BD5", showValue=True),
            )

    pie_start = cat_start
    pie_col = 8

    ws.cell(row=pie_start, column=pie_col, value="Priority Breakdown").font = _SECTION_FONT

    ws.cell(row=pie_start + 1, column=pie_col, value="Priority").font = _HEADER_FONT
    ws.cell(row=pie_start + 1, column=pie_col).fill = _HEADER_FILL
    ws.cell(row=pie_start + 1, column=pie_col).alignment = _CENTER
    ws.cell(row=pie_start + 1, column=pie_col).border = _THIN_BORDER
    ws.cell(row=pie_start + 1, column=pie_col + 1, value="Count").font = _HEADER_FONT
    ws.cell(row=pie_start + 1, column=pie_col + 1).fill = _HEADER_FILL
    ws.cell(row=pie_start + 1, column=pie_col + 1).alignment = _CENTER
    ws.cell(row=pie_start + 1, column=pie_col + 1).border = _THIN_BORDER

    prio_row = pie_start + 2
    sorted_prios = sorted(priorities.items(), key=lambda x: x[1], reverse=True)
    for i, (name, count) in enumerate(sorted_prios):
        r = prio_row + i
        ws.cell(row=r, column=pie_col, value=name)
        ws.cell(row=r, column=pie_col + 1, value=count)
        _style_body(ws, r, pie_col, pie_col + 1)

    if priorities:
        pie = PieChart()
        pie.title = "Snags by Priority"
        pie.style = 10
        data_ref = Reference(ws, min_col=pie_col + 1, min_row=pie_start + 1, max_row=prio_row + len(priorities) - 1)
        cats_ref = Reference(ws, min_col=pie_col, min_row=prio_row, max_row=prio_row + len(priorities) - 1)
        pie.add_data(data_ref, titles_from_data=True)
        pie.set_categories(cats_ref)
        ws.add_chart(pie, f"I{pie_start}")

    vend_start = max(cat_row + len(categories) + 3, prio_row + len(priorities) + 2)

    ws.merge_cells(f"A{vend_start}:E{vend_start}")
    ws.cell(row=vend_start, column=1, value="Vendor Distribution").font = _SECTION_FONT

    ws.cell(row=vend_start + 1, column=1, value="Vendor").font = _HEADER_FONT
    ws.cell(row=vend_start + 1, column=1).fill = _HEADER_FILL
    ws.cell(row=vend_start + 1, column=1).alignment = _CENTER
    ws.cell(row=vend_start + 1, column=1).border = _THIN_BORDER
    ws.cell(row=vend_start + 1, column=2, value="Count").font = _HEADER_FONT
    ws.cell(row=vend_start + 1, column=2).fill = _HEADER_FILL
    ws.cell(row=vend_start + 1, column=2).alignment = _CENTER
    ws.cell(row=vend_start + 1, column=2).border = _THIN_BORDER

    sorted_vendors = sorted(vendors.items(), key=lambda x: x[1], reverse=True)
    for i, (name, count) in enumerate(sorted_vendors):
        r = vend_start + 2 + i
        ws.cell(row=r, column=1, value=name)
        ws.cell(row=r, column=2, value=count)
        _style_body(ws, r, 1, 2)

    ws.column_dimensions["A"].width = 26
    ws.column_dimensions["B"].width = 14
    ws.column_dimensions["C"].width = 10
    ws.column_dimensions["D"].width = 14
    ws.column_dimensions["E"].width = 14

    logger.info("Summary dashboard sheet added with %d categories, %d priorities, %d vendors",
                len(categories), len(priorities), len(vendors))
