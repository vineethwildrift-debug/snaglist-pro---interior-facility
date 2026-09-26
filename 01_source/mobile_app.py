from __future__ import annotations

import os
from pathlib import Path

from flask import Flask, render_template_string, request, send_file
from openpyxl import Workbook, load_workbook

from desktop_app import _apply_optional_spelling_correction, _validate_export_workbook
from desktop_app import _build_export_rows, _decode_archive_text, _read_checklist_rows, _archive_text_candidates


HTML_TEMPLATE = """
<!doctype html>
<html lang="en">
<head>
  <meta charset="utf-8">
  <meta name="viewport" content="width=device-width, initial-scale=1">
  <title>Snaglist Pro Mobile</title>
  <style>
    body { font-family: Arial, sans-serif; margin: 0; background: #f1f5f9; color: #0f172a; }
    .shell { max-width: 720px; margin: 0 auto; padding: 24px 18px 60px; }
    .card { background: #fff; border-radius: 14px; padding: 20px; box-shadow: 0 10px 25px rgba(15, 23, 42, 0.08); }
    h1 { margin-top: 0; }
    label { display: block; font-weight: bold; margin: 18px 0 8px; }
    input[type="file"], button { width: 100%; box-sizing: border-box; }
    input[type="file"] { padding: 10px; border: 1px solid #cbd5e1; border-radius: 8px; }
    button { background: #2563eb; color: white; border: none; border-radius: 10px; padding: 14px; font-weight: bold; margin-top: 18px; }
    .muted { color: #475569; font-size: 0.95rem; }
    .note { background: #eff6ff; border-left: 4px solid #2563eb; padding: 12px 14px; margin-top: 18px; border-radius: 8px; }
  </style>
</head>
<body>
  <div class="shell">
    <div class="card">
      <h1>Snaglist Pro</h1>
      <p class="muted">Mobile-first WhatsApp ZIP and checklist export.</p>
      <form method="post" enctype="multipart/form-data" action="/generate">
        <label for="zip">WhatsApp ZIP export</label>
        <input id="zip" name="zip" type="file" accept=".zip" required>

        <label for="checklist">Checklist file</label>
        <input id="checklist" name="checklist" type="file" accept=".csv,.xlsx" required>

        <label for="correct_spelling">
          <input id="correct_spelling" name="correct_spelling" type="checkbox" checked> Correct spelling in report
        </label>

        <button type="submit">Generate workbook</button>
      </form>
      <div class="note">Validated export flow keeps the Summary sheet and audit notes in place before download.</div>
    </div>
  </div>
</body>
</html>
"""


def create_mobile_app():
    app = Flask(__name__)

    @app.get("/")
    def index():
        return render_template_string(HTML_TEMPLATE)

    @app.post("/generate")
    def generate_route():
        zip_file = request.files.get("zip")
        checklist_file = request.files.get("checklist")
        correct_spelling = request.form.get("correct_spelling") in {"on", "true", "1", "yes"}

        if not zip_file or not checklist_file:
            return "ZIP export and checklist are required.", 400

        temp_dir = Path(os.environ.get("LOCALAPPDATA", str(Path.home()))) / "SnaglistProMobile" / "uploads"
        temp_dir.mkdir(parents=True, exist_ok=True)

        zip_path = temp_dir / f"upload_{abs(hash(zip_file.filename))}.zip"
        checklist_path = temp_dir / f"upload_{abs(hash(checklist_file.filename))}.csv"

        zip_file.save(zip_path)
        checklist_file.save(checklist_path)

        try:
            output_path = generate_report_from_uploads(str(zip_path), str(checklist_path), correct_spelling=correct_spelling)
        except ValueError as exc:
            return str(exc), 400

        return send_file(output_path, as_attachment=True, download_name=Path(output_path).name)

    return app


def generate_report_from_uploads(zip_path: str, checklist_path: str, correct_spelling: bool = True):
    if not zip_path or not checklist_path:
        raise ValueError("ZIP export and checklist are required.")

    zip_file = Path(zip_path)
    checklist_file = Path(checklist_path)

    if not zip_file.exists():
        raise ValueError("ZIP export was not found.")
    if not checklist_file.exists():
        raise ValueError("Checklist file was not found.")

    import zipfile

    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Snaglist"

    rows = []
    with zipfile.ZipFile(zip_file, "r") as archive:
        text_files = _archive_text_candidates(archive)
        parsed_rows = []
        for name in text_files:
            try:
                raw_text = _decode_archive_text(archive.read(name))
                parsed_rows.extend(__import__("snaglist_pro.parse_export", fromlist=["parse_whatsapp_text"]).parse_whatsapp_text(raw_text))
            except Exception:
                continue

        checklist_rows = _read_checklist_rows(str(checklist_file))
        rows = _build_export_rows(parsed_rows, checklist_rows)

    if not rows:
        raise ValueError("No valid data rows were found in the ZIP or checklist.")

    sheet.append([
        "Date", "Location", "Item", "Description", "Category", "Status", "Priority",
        "Source", "Media Files", "Notes",
    ])

    for row in rows:
        normalized = {}
        for key in [
            "Date", "Location", "Item", "Description", "Category", "Status", "Priority",
            "Source", "Media Files", "Notes",
        ]:
            value = row.get(key, "")
            if correct_spelling and key in {"Description", "Item", "Location", "Notes"}:
                value = _apply_optional_spelling_correction(value, True)
            normalized[key] = value
        sheet.append([normalized.get(column, "") for column in [
            "Date", "Location", "Item", "Description", "Category", "Status", "Priority",
            "Source", "Media Files", "Notes",
        ]])

    summary = workbook.create_sheet("Summary")
    summary.append(["Metric", "Value"])
    summary.append(["Total Items", len(rows)])
    summary.append(["Open", sum(1 for row in rows if str(row.get("Status", "")).strip().lower() in {"open", "new"})])
    summary.append(["In Progress", sum(1 for row in rows if str(row.get("Status", "")).strip().lower() in {"in progress", "pending"})])
    summary.append(["Resolved", sum(1 for row in rows if str(row.get("Status", "")).strip().lower() in {"closed", "resolved", "done"})])

    issues = _validate_export_workbook(workbook)
    if issues:
        raise ValueError("Workbook validation failed: " + "; ".join(issues))

    output_dir = Path(os.environ.get("LOCALAPPDATA", str(Path.home()))) / "SnaglistProMobile" / "exports"
    output_dir.mkdir(parents=True, exist_ok=True)
    output_path = output_dir / f"snaglist_mobile_{abs(hash(str(zip_path) + str(checklist_path) + str(Path().stat().st_mtime)))}.xlsx"
    workbook.save(output_path)
    return str(output_path)


if __name__ == "__main__":
    create_mobile_app().run(host="0.0.0.0", port=5000, debug=False)
