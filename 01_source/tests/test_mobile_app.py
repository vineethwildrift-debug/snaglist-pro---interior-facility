from pathlib import Path

import pytest
from openpyxl import load_workbook

from mobile_app import create_mobile_app, generate_report_from_uploads


def _write_sample_files(tmp_path):
    zip_path = tmp_path / "sample.zip"
    checklist_path = tmp_path / "checklist.csv"

    checklist_path.write_text(
        "Location,Item,Description,Status,Priority\nLobby,Light,Emeregency allignment,Open,High\n",
        encoding="utf-8",
    )

    import zipfile

    with zipfile.ZipFile(zip_path, "w") as archive:
        archive.writestr("WhatsApp Chat Export.txt", "20/09/2026, 10:00 - Inspector: Door handle loose")

    return zip_path, checklist_path


def test_mobile_app_index_page_loads():
    app = create_mobile_app()

    with app.test_client() as client:
        response = client.get("/")
        assert response.status_code == 200
        assert b"Snaglist Pro" in response.data


def test_mobile_app_generates_valid_workbook(tmp_path):
    zip_path, checklist_path = _write_sample_files(tmp_path)

    output_path = generate_report_from_uploads(str(zip_path), str(checklist_path))

    assert output_path
    assert Path(output_path).exists()

    workbook = load_workbook(output_path)
    assert "Summary" in workbook.sheetnames
    assert "Snaglist" in workbook.sheetnames


def test_mobile_app_requires_zip_and_checklist():
    with pytest.raises(ValueError):
        generate_report_from_uploads("", "")
