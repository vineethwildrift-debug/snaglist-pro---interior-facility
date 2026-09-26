from pathlib import Path


ROOT = Path(__file__).resolve().parent.parent
APP_LAUNCHER = ROOT / "app_launcher.py"


def test_launcher_uses_native_desktop_entrypoint():
    text = APP_LAUNCHER.read_text(encoding="utf-8")
    assert "run_desktop_app" in text or "Tk" in text


def test_native_desktop_module_exists():
    native_module = ROOT / "desktop_app.py"
    assert native_module.exists()


def test_desktop_shell_uses_modern_theme():
    native_module = ROOT / "desktop_app.py"
    text = native_module.read_text(encoding="utf-8")
    assert "0f172a" in text
    assert "94a3b8" in text
    assert "modern" in text.lower()


def test_desktop_app_escapes_multiword_font_family():
    native_module = ROOT / "desktop_app.py"
    text = native_module.read_text(encoding="utf-8")
    assert 'option_add("*Font", "{Segoe UI} 10")' in text


def test_desktop_app_has_download_action():
    native_module = ROOT / "desktop_app.py"
    text = native_module.read_text(encoding="utf-8")
    assert "Download Excel/Sheets" in text
    assert "download_output" in text
    assert ".xlsx" in text


def test_desktop_generation_uses_canonical_pipeline_output():
    native_module = ROOT / "desktop_app.py"
    text = native_module.read_text(encoding="utf-8")
    assert "SnaglistPipeline" in text
    assert 'result.get("excel_path")' in text


def test_desktop_app_has_optional_spelling_toggle_and_validation_helper():
    import desktop_app

    assert desktop_app._apply_optional_spelling_correction("Emeregency allignment", True) == "Emergency alignment"
    assert desktop_app._apply_optional_spelling_correction("Emeregency allignment", False) == "Emeregency allignment"

    workbook = desktop_app.Workbook()
    sheet = workbook.active
    sheet.title = "Snaglist"
    sheet.append(["Date", "Location", "Item", "Description"])
    sheet.append(["2026-09-20", "Lobby", "Lighting", "Emeregency allignment"])

    issues = desktop_app._validate_export_workbook(workbook)
    assert "Missing Summary sheet" in issues

    summary = workbook.create_sheet("Summary")
    summary.append(["Status", "Value"])
    summary.append(["Open", 1])

    assert desktop_app._validate_export_workbook(workbook) == []


def test_report_spelling_correction_preserves_technical_terms():
    from snaglist_pro.pipeline import correct_report_text

    value = "Emeregency allignment near AHU and VTPN DB"

    assert correct_report_text(value) == "Emergency alignment near AHU and VTPN DB"


def test_report_spelling_correction_returns_change_log():
    from snaglist_pro.pipeline import correct_report_text_with_changes

    corrected, changes = correct_report_text_with_changes("Diffueser and Landscapping")

    assert corrected == "Diffuser and Landscaping"
    assert changes == [
        ("Diffueser", "Diffuser"),
        ("Landscapping", "Landscaping"),
    ]


def test_download_name_sanitizer_does_not_duplicate_xlsx_suffix():
    import desktop_app

    for name in [
        "snaglist_report_20260920_165128.xlsx",
        "snaglist_report_20260920_165128.xls",
        "snaglist_report_20260920_165128.csv",
    ]:
        safe_name = desktop_app._suggest_download_filename(name)
        assert safe_name.lower().endswith(".xlsx")
        assert not safe_name.lower().endswith(".xlsxx.xlsx")
        assert safe_name.lower().count(".xlsx") == 1


def test_project_export_uses_snag_columns_not_raw_chat_headers():
    import desktop_app

    rows = desktop_app._build_export_rows(
        [{"date": "2026-09-20", "time": "10:00", "sender": "Alice", "message": "Door handle loose", "has_media": False, "media_files": []}],
        [],
    )

    assert rows
    assert "Location" in rows[0]
    assert "Description" in rows[0]
    assert "Status" in rows[0]
    assert "Priority" in rows[0]
    assert "Date" in rows[0]
    assert "Raw Sender" not in rows[0]


def test_checklist_rows_are_included_in_export_data():
    import desktop_app

    rows = desktop_app._build_export_rows(
        [],
        [
            {"Location": "Living Room", "Item": "Door Handle", "Description": "Loose handle", "Status": "Open", "Priority": "High", "Category": "Joinery"},
            {"Location": "Kitchen", "Item": "Cabinet", "Description": "Misaligned panel", "Status": "In Progress", "Priority": "Medium", "Category": "Carpentry"},
        ],
    )

    assert len(rows) == 2
    assert rows[0]["Location"] == "Living Room"
    assert rows[0]["Description"] == "Loose handle"
    assert rows[0]["Status"] == "Open"
    assert rows[1]["Location"] == "Kitchen"


def test_checklist_reader_finds_template_header_row_and_maps_questions(tmp_path):
    import desktop_app
    from openpyxl import Workbook

    checklist = tmp_path / "template.xlsx"
    workbook = Workbook()
    sheet = workbook.active
    sheet.append([])
    sheet.append(["Index"])
    sheet.append([])
    sheet.append([])
    sheet.append(["Slno", "Facility Name", "Catagory", "Check points", "Priority"])
    sheet.append([1, "", "Electrical", "Is all lights installed?", "Medium"])
    workbook.save(checklist)

    rows = desktop_app._read_checklist_rows(str(checklist))

    assert len(rows) == 1
    assert rows[0]["Slno"] == 1
    assert rows[0]["Catagory"] == "Electrical"
    assert rows[0]["Check points"] == "Is all lights installed?"


def test_image_size_selector_has_200x200_option():
    import desktop_app

    assert "200x200" in desktop_app.IMAGE_SIZE_OPTIONS
    assert desktop_app._normalize_image_size("200x200") == 200


def test_apply_image_to_cell_resizes_for_excel_sheet():
    import desktop_app
    from pathlib import Path

    from PIL import Image
    from openpyxl import Workbook

    png = Path("tmp_test_image.png")
    Image.new("RGB", (40, 40), color="blue").save(png)

    try:
        wb = Workbook()
        ws = wb.active
        ws.title = "Snaglist"
        ws["A1"] = "Sample"
        desktop_app._attach_image_to_cell(ws, 1, "A", str(png), 200)
        assert ws.row_dimensions[1].height > 0
        assert ws.column_dimensions["A"].width > 0
        assert len(ws._images) == 1
    finally:
        if png.exists():
            png.unlink()


def test_fit_inside_cell_preserves_aspect_ratio():
    import desktop_app

    assert desktop_app._fit_image_dimensions(800, 400, 200) == (200, 100)
    assert desktop_app._fit_image_dimensions(400, 800, 200) == (100, 200)


def test_decode_whatsapp_text_supports_utf16_exports():
    import desktop_app

    content = "20/09/2026, 10:00 - Inspector: Door handle loose"

    assert desktop_app._decode_archive_text(content.encode("utf-16")) == content


def test_parser_supports_bracketed_whatsapp_exports():
    from snaglist_pro.parse_export import parse_whatsapp_text

    rows = parse_whatsapp_text("[20/09/2026, 10:00] Inspector: Door handle loose")

    assert len(rows) == 1
    assert rows[0]["sender"] == "Inspector"
    assert rows[0]["message"] == "Door handle loose"


def test_archive_text_candidates_include_extensionless_text_files():
    import io
    import zipfile
    import desktop_app

    archive_buffer = io.BytesIO()
    with zipfile.ZipFile(archive_buffer, "w") as archive:
        archive.writestr("WhatsApp Chat Export", "20/09/2026, 10:00 - Inspector: Door handle loose")
        archive.writestr("photo.jpg", b"\xff\xd8\xff\xe0\x00\x10JFIF")

    with zipfile.ZipFile(io.BytesIO(archive_buffer.getvalue())) as archive:
        assert desktop_app._archive_text_candidates(archive) == ["WhatsApp Chat Export"]


def test_extract_zip_clears_previous_project_files(tmp_path):
    import zipfile
    from snaglist_pro.pipeline import extract_zip

    archive_path = tmp_path / "navigator.zip"
    extract_dir = tmp_path / "extracted"
    extract_dir.mkdir()
    (extract_dir / "old-project.txt").write_text("old project", encoding="utf-8")

    with zipfile.ZipFile(archive_path, "w") as archive:
        archive.writestr("chat.txt", "new project chat")

    extract_zip(str(archive_path), str(extract_dir))

    assert not (extract_dir / "old-project.txt").exists()
    assert (extract_dir / "chat.txt").exists()


def test_project_metadata_is_not_bound_to_etv_defaults():
    from snaglist_pro.pipeline import detect_project_floor, slugify_project

    assert detect_project_floor("Navigater - 6F - Rodeofx India") == "6F"
    assert slugify_project("Navigater - 6F - Rodeofx India") == "navigater__6f__rodeofx_india"


def test_summary_dashboard_uses_cell_references_for_percentages():
    from openpyxl import Workbook
    from snaglist_pro.summary_dashboard import add_summary_sheet

    workbook = Workbook()
    add_summary_sheet(
        workbook,
        {"total": 10, "matched": 4, "unmatched": 6, "unmatched_snags": 2, "open": 10, "closed": 0,
         "categories": {"Electrical": 10}, "priorities": {"Medium": 10}, "vendors": {}},
        data_sheet="Snaglist",
    )

    summary = workbook["Summary"]

    assert summary["B9"].value == "=B6/(B6+B8)"
    assert summary["C14"].value == "=B14/$B$5"


def test_summary_sheet_open_closed_counts_reflect_actual_data(sample_checklist_path, output_dir):
    import tempfile
    import zipfile
    import os
    from PIL import Image
    from snaglist_pro.pipeline import SnaglistPipeline
    from openpyxl import load_workbook

    tmp_dir = tempfile.mkdtemp()
    chat_text = (
        "26/07/2026, 10:30:00 - Inspector: Window to be rectified\n"
        "26/07/2026, 10:31:00 - Inspector: IMG-20260726-WA0001.jpg (file attached)\n"
        "26/07/2026, 10:32:00 - Inspector: Door signage fixed\n"
        "26/07/2026, 10:33:00 - Inspector: IMG-20260726-WA0002.jpg (file attached)\n"
        "26/07/2026, 10:34:00 - Inspector: Tiles damaged and to be replaced\n"
        "26/07/2026, 10:35:00 - Inspector: IMG-20260726-WA0003.jpg (file attached)\n"
    )
    extract_dir = os.path.join(tmp_dir, "extracted")
    os.makedirs(extract_dir)
    img_dir = os.path.join(extract_dir, "images")
    os.makedirs(img_dir)
    for i in range(3):
        img = Image.new("RGB", (100, 100), color=(50 * (i + 1), 30 * (i + 1), 20 * (i + 1)))
        img.save(os.path.join(img_dir, f"IMG-20260726-WA00{i + 1}.jpg"), "JPEG")
    with open(os.path.join(extract_dir, "_chat.txt"), "w", encoding="utf-8") as f:
        f.write(chat_text)
    zip_path = os.path.join(tmp_dir, "TestProject.zip")
    with zipfile.ZipFile(zip_path, "w", zipfile.ZIP_DEFLATED) as zf:
        for root, _dirs, files in os.walk(extract_dir):
            for fname in files:
                fpath = os.path.join(root, fname)
                arcname = os.path.relpath(fpath, extract_dir)
                zf.write(fpath, arcname)

    pipeline = SnaglistPipeline()
    result = pipeline.run(
        zip_path=zip_path,
        checklist_path=sample_checklist_path,
        output_dir=output_dir,
    )
    assert result["total_snags"] > 0

    wb = load_workbook(result["excel_path"])
    data_ws = None
    for ws in wb.worksheets:
        if ws.title not in ("Summary", "Correction Notes"):
            data_ws = ws
            break
    assert data_ws is not None

    actual_open = 0
    actual_closed = 0
    for row in data_ws.iter_rows(min_row=8, values_only=True):
        status = row[8]
        if status == "Open":
            actual_open += 1
        elif status == "Closed":
            actual_closed += 1

    summary = wb["Summary"]
    assert summary["B10"].value == actual_open
    assert summary["B11"].value == actual_closed
    assert summary["B10"].value + summary["B11"].value == summary["B5"].value
    assert actual_closed > 0


def test_image_references_without_media_files_remain_as_snags():
    from snaglist_pro.pipeline import pair_images_with_descriptions

    messages = [
        {"date": "2026-09-21", "sender": "Inspector", "message": "<image omitted> Door alignment needs repair"},
        {"date": "2026-09-21", "sender": "Inspector", "message": "<image omitted> Paint touch-up required"},
    ]

    snags = pair_images_with_descriptions(messages, {})

    assert len(snags) == 2
    assert snags[0]["description"] == "Door alignment needs repair"
    assert snags[0]["image_path"] == ""
    assert snags[0]["media_missing"] is True


def test_pipeline_parser_supports_bracketed_whatsapp_headers():
    from snaglist_pro.pipeline import parse_chat_messages

    messages = parse_chat_messages(
        "[9/10/26, 10:50:42 AM] Imtiyaz Indiqube: [Image] Door alignment needs repair"
    )

    assert len(messages) == 1
    assert messages[0]["sender"] == "Imtiyaz Indiqube"
    assert messages[0]["message"] == "[Image] Door alignment needs repair"


def test_checklist_can_generate_when_zip_has_no_chat_text():
    import zipfile
    from types import SimpleNamespace

    import desktop_app

    checklist = Path("tmp_checklist.csv")
    archive_path = Path("tmp_media_only.zip")
    checklist.write_text("Location,Item,Description\nKitchen,Sink,Loose tap\n", encoding="utf-8")
    with zipfile.ZipFile(archive_path, "w") as archive:
        archive.writestr("README", "No chat export included")

    try:
        app = desktop_app.SnaglistDesktopApp.__new__(desktop_app.SnaglistDesktopApp)
        app.selected_checklist = str(checklist)
        app.root = SimpleNamespace(update_idletasks=lambda: None)
        app.status_var = SimpleNamespace(set=lambda value: None)
        app.progress_var = SimpleNamespace(set=lambda value: None)
        app.image_size_var = SimpleNamespace(get=lambda: "200x200")
        app.fit_inside_cell_var = SimpleNamespace(get=lambda: True)
        app.append_log = lambda message: None

        file_bytes, suggested_name = app._generate_output_from_zip(str(archive_path))
        assert file_bytes
        assert isinstance(file_bytes, bytes)
        assert len(file_bytes) > 0
        assert suggested_name and suggested_name.endswith(".xlsx")
        tmp = Path("tmp_generated.xlsx")
        try:
            tmp.write_bytes(file_bytes)
            assert tmp.exists()
        finally:
            tmp.unlink(missing_ok=True)
    finally:
        checklist.unlink(missing_ok=True)
        archive_path.unlink(missing_ok=True)


def test_unrecognized_chat_text_can_still_be_exported():
    import desktop_app

    rows = desktop_app._fallback_rows_from_text("Inspection note without WhatsApp date\nSecond note")

    assert len(rows) == 2
    assert rows[0]["Description"] == "Inspection note without WhatsApp date"
    assert rows[0]["Source"] == "ZIP text"


def test_media_only_zip_can_create_export_rows():
    import desktop_app

    rows = desktop_app._fallback_rows_from_media(["IMG-001.jpg", "IMG-002.jpg"])

    assert len(rows) == 2
    assert rows[0]["Item"] == "IMG-001.jpg"
    assert rows[0]["Source"] == "ZIP media"


def test_report_image_selection_limits_embedded_images():
    import desktop_app

    names = [f"IMG-{index:03d}.jpg" for index in range(25)]

    assert len(desktop_app._report_image_names(names)) == 10
    assert desktop_app._report_image_names(names) == names[:10]


def test_prepare_image_for_excel_reduces_physical_dimensions():
    from PIL import Image
    import desktop_app

    source = Path("tmp_large_image.jpg")
    target = Path("tmp_resized_image.jpg")
    Image.new("RGB", (2400, 1600), color="red").save(source, quality=95)

    try:
        desktop_app._prepare_image_for_excel(str(source), str(target), 200, True)
        with Image.open(target) as image:
            assert image.size == (200, 133)
    finally:
        source.unlink(missing_ok=True)
        target.unlink(missing_ok=True)
