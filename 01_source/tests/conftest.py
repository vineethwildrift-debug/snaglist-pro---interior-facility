import os
import sys
import tempfile
import zipfile
from pathlib import Path
from datetime import datetime

import pytest

sys.path.insert(0, str(Path(__file__).resolve().parent.parent))


SAMPLE_CHAT_TEXT = """
26/07/2026, 10:30:00 - Nasar Indiqube: Meeting cabin Fascad window to be rectified
26/07/2026, 10:31:00 - Nasar Indiqube: IMG-20260726-WA0001.jpg (file attached)
26/07/2026, 10:32:00 - Nasar Indiqube: Door push-pull signage to be installed
26/07/2026, 10:33:00 - Nasar Indiqube: IMG-20260726-WA0002.jpg (file attached)
26/07/2026, 10:34:00 - vineeth bisaa: <Media omitted>
26/07/2026, 10:35:00 - vineeth bisaa: Damaged tiles to be replaced - MR cabin 1
26/07/2026, 10:36:00 - Nasar Indiqube: created group "Test Project - 5F"
26/07/2026, 10:37:00 - Nasar Indiqube: Ok
26/07/2026, 10:38:00 - Nasar Indiqube: IMG-20260726-WA0003.jpg (file attached)
26/07/2026, 10:39:00 - vineeth bisaa: Proper support to be provided for AC duct
"""


@pytest.fixture
def sample_chat_text() -> str:
    return SAMPLE_CHAT_TEXT


@pytest.fixture
def sample_zip_path() -> str:
    tmp_dir = tempfile.mkdtemp()
    zip_path = os.path.join(tmp_dir, "WhatsApp Chat with Test.zip")
    extract_dir = os.path.join(tmp_dir, "extracted")
    os.makedirs(extract_dir)

    txt_path = os.path.join(extract_dir, "_chat.txt")
    with open(txt_path, "w", encoding="utf-8") as f:
        f.write(SAMPLE_CHAT_TEXT)

    img_dir = os.path.join(extract_dir, "images")
    os.makedirs(img_dir)

    from PIL import Image
    for i in range(1, 4):
        img = Image.new("RGB", (100, 100), color=(i * 50, i * 30, i * 20))
        img.save(os.path.join(img_dir, f"IMG-20260726-WA000{i}.jpg"), "JPEG")

    with zipfile.ZipFile(zip_path, "w", zipfile.ZIP_DEFLATED) as zf:
        for root, _dirs, files in os.walk(extract_dir):
            for fname in files:
                fpath = os.path.join(root, fname)
                arcname = os.path.relpath(fpath, extract_dir)
                zf.write(fpath, arcname)

    return zip_path


@pytest.fixture
def sample_checklist_path() -> str:
    import openpyxl

    tmp_dir = tempfile.mkdtemp()
    xlsx_path = os.path.join(tmp_dir, "checklist.xlsx")

    wb = openpyxl.Workbook()
    ws = wb.active
    ws["A1"] = "Slno"
    ws["G1"] = "Catagory"
    ws["H1"] = "Check points"
    ws["K1"] = "Priority"

    items = [
        (1, "Interior", "Window to be rectified", "Medium"),
        (2, "Interior", "Door signage to be installed", "Medium"),
        (3, "Civil", "Damaged tiles to be replaced", "High"),
        (4, "HVAC", "AC duct support to be provided", "Medium"),
    ]
    for slno, cat, check, pri in items:
        row = slno + 1
        ws.cell(row=row, column=1, value=slno)
        ws.cell(row=row, column=7, value=cat)
        ws.cell(row=row, column=8, value=check)
        ws.cell(row=row, column=11, value=pri)

    wb.save(xlsx_path)
    return xlsx_path


@pytest.fixture
def output_dir() -> str:
    return tempfile.mkdtemp()
