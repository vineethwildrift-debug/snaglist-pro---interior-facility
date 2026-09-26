import os
import json
import pytest
from snaglist_pro.pipeline import (
    detect_project_name,
    slugify_project,
    parse_chat_messages,
    pair_images_with_descriptions,
    deduplicate_snags,
    match_to_checklist,
    determine_priority,
    get_vendor,
    load_checklist,
    SnaglistPipeline,
)


class TestDetectProjectName:
    def test_from_zip_filename(self):
        name = detect_project_name(None, "WhatsApp Chat with ETV 6F.zip")
        assert "ETV 6F" in name

    def test_from_created_group_message(self):
        text = 'Nasar created group "Test Project"'
        name = detect_project_name(text, "some.zip")
        assert name == "Test Project"

    def test_created_group_takes_priority(self):
        text = 'created group "Group Name"'
        name = detect_project_name(text, "WhatsApp Chat with Other.zip")
        assert name == "Group Name"

    def test_fallback_to_unknown(self):
        name = detect_project_name(None, "")
        assert name == "Unknown Project"


class TestSlugifyProject:
    def test_basic(self):
        assert slugify_project("ETV 6F") == "etv_6f"

    def test_removes_special_chars(self):
        assert slugify_project("Test! Project @2024") == "test_project_2024"

    def test_fallback(self):
        assert slugify_project("!!!") != ""


class TestParseChatMessages:
    def test_parses_messages(self, sample_chat_text):
        messages = parse_chat_messages(sample_chat_text)
        assert len(messages) > 0

    def test_parses_sender(self, sample_chat_text):
        messages = parse_chat_messages(sample_chat_text)
        assert all(m.get("sender") for m in messages)


class TestDeduplicateSnags:
    def test_removes_duplicates(self):
        snags = [
            {"image_filename": "a.jpg", "description": "test desc"},
            {"image_filename": "a.jpg", "description": "test desc"},
            {"image_filename": "b.jpg", "description": "other desc"},
        ]
        result = deduplicate_snags(snags)
        assert len(result) == 2

    def test_no_duplicates(self):
        snags = [
            {"image_filename": "a.jpg", "description": "desc1"},
            {"image_filename": "b.jpg", "description": "desc2"},
        ]
        result = deduplicate_snags(snags)
        assert len(result) == 2

    def test_empty_input(self):
        assert deduplicate_snags([]) == []


class TestMatchToChecklist:
    def test_matches_by_category_and_jaccard(self):
        checklist = [
            {"slno": 1, "category": "Electrical", "check_points": "Light not working", "priority": "Medium"},
            {"slno": 2, "category": "Interior", "check_points": "Window to be rectified", "priority": "Medium"},
        ]
        item, score = match_to_checklist("Window damaged", "Interior", checklist)
        assert item is not None
        assert "window" in item["check_points"].lower()

    def test_no_match_returns_none(self):
        checklist = [
            {"slno": 1, "category": "Electrical", "check_points": "Light switch", "priority": "Medium"},
        ]
        item, score = match_to_checklist("ZZZZZ no match possible here", "Interior", checklist)
        assert item is None

    def test_empty_description(self):
        item, score = match_to_checklist("", "Electrical", [{"slno": 1, "category": "Electrical", "check_points": "x", "priority": "M"}])
        assert item is None


class TestDeterminePriority:
    def test_high_from_keyword(self):
        assert determine_priority("urgent leak detected", None) == "High"
        assert determine_priority("fire hazard", None) == "High"

    def test_default_from_matched_item(self):
        matched = {"priority": "Low"}
        assert determine_priority("paint touch up needed", matched) == "Low"

    def test_default_priority(self):
        assert determine_priority("minor issue", None) == "Medium"


class TestGetVendor:
    def test_electrical(self):
        assert get_vendor("Electrical") == "Ever green"

    def test_hvac(self):
        assert get_vendor("Hvac") == "Apel (Carrier)"

    def test_unknown(self):
        assert get_vendor("Unknown Category") == ""

    def test_empty(self):
        assert get_vendor("") == ""


class TestLoadChecklist:
    def test_skips_header_subheader_rows(self, tmp_path):
        import openpyxl

        xlsx_path = str(tmp_path / "checklist.xlsx")
        wb = openpyxl.Workbook()
        ws = wb.active
        ws["A1"] = "Slno"
        ws["G1"] = "Catagory"
        ws["H1"] = "Check points"
        ws["K1"] = "Priority"
        ws["A2"] = "Snagpoints"
        ws["G2"] = "Catagory"
        ws["H2"] = "Check Points"
        ws["K2"] = "Priority"
        ws["A3"] = 1
        ws["G3"] = "Interior"
        ws["H3"] = "Window to be rectified"
        ws["K3"] = "Medium"
        ws["A4"] = 2
        ws["G4"] = "Civil"
        ws["H4"] = "Damaged tiles"
        ws["K4"] = "High"
        wb.save(xlsx_path)

        items = load_checklist(xlsx_path)
        assert len(items) == 2
        categories = [i["category"] for i in items]
        assert "Catagory" not in categories
        assert categories == ["Interior", "Civil"]


class TestSnaglistPipeline:
    def test_pipeline_runs(self, sample_zip_path, sample_checklist_path, output_dir):
        pipeline = SnaglistPipeline()
        result = pipeline.run(
            zip_path=sample_zip_path,
            checklist_path=sample_checklist_path,
            output_dir=output_dir,
        )
        assert result["total_snags"] > 0
        assert os.path.exists(result["excel_path"])
        assert result["project"] is not None

    def test_pipeline_generates_json(self, sample_zip_path, sample_checklist_path, output_dir):
        pipeline = SnaglistPipeline()
        result = pipeline.run(
            zip_path=sample_zip_path,
            checklist_path=sample_checklist_path,
            output_dir=output_dir,
        )
        data_path = os.path.join(output_dir, "snag_data.json")
        assert os.path.exists(data_path)
        with open(data_path) as f:
            data = json.load(f)
        assert "snags" in data
        assert data["total_snags"] > 0
