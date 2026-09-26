import csv
import json
import logging
import os
import sys
import tkinter as tk
import zipfile
from datetime import datetime
from pathlib import Path
from tkinter import filedialog, messagebox, ttk

from PIL import Image
from openpyxl import Workbook, load_workbook
from openpyxl.drawing.image import Image as XLImage

from snaglist_pro.parse_export import parse_whatsapp_text
from snaglist_pro.config import settings


WINDOW_TITLE = "Snaglist Pro"
DEFAULT_USERNAME = os.environ.get("SNAGLIST_APP_USERNAME", "")
DEFAULT_PASSWORD = os.environ.get("SNAGLIST_APP_PASSWORD", "")


def _get_output_dir() -> Path:
    configured = (settings.project_output_dir or "").strip()
    if configured:
        return Path(configured)
    return Path(os.environ.get("LOCALAPPDATA", Path.home())) / "SnaglistPro" / "exports"


def _setup_file_logging():
    """Add a file handler that logs errors to AppData/snaglist_error.log."""
    try:
        app_dir = Path(os.environ.get("LOCALAPPDATA", Path.home())) / "SnaglistPro"
        app_dir.mkdir(parents=True, exist_ok=True)
        log_file = app_dir / "snaglist_error.log"
        handler = logging.FileHandler(str(log_file), encoding="utf-8")
        handler.setLevel(logging.ERROR)
        formatter = logging.Formatter("%(asctime)s [%(levelname)s] %(message)s")
        handler.setFormatter(formatter)
        logging.getLogger().addHandler(handler)
    except Exception:
        pass


def _safe_font(size, weight="normal"):
    family = "Segoe UI"
    if weight and weight.lower() == "bold":
        return (family, size, "bold")
    return (family, size)


def _suggest_download_filename(filename: str) -> str:
    """Normalize download names so .xlsx is added once and not duplicated."""
    if not filename:
        return "snaglist_report.xlsx"

    clean_name = Path(filename).name
    lowered = clean_name.lower()
    for suffix in (".xlsx", ".xlsm", ".xls", ".csv"):
        if lowered.endswith(suffix):
            clean_name = clean_name[: -len(suffix)]
            break

    clean_name = clean_name.strip() or "snaglist_report"
    return f"{clean_name}.xlsx"


EXPORT_COLUMNS = [
    "Date",
    "Location",
    "Item",
    "Description",
    "Category",
    "Status",
    "Priority",
    "Source",
    "Media Files",
    "Notes",
]

IMAGE_SIZE_OPTIONS = ["200x200", "150x150", "100x100", "80x80"]
FIT_INSIDE_CELL_LABEL = "Fit inside cell"
SPELLING_CORRECTION_LABEL = "Correct spelling in report"


def _apply_optional_spelling_correction(value, enabled=True):
    if not enabled:
        return value
    if value is None:
        return value
    text = str(value)
    if not text:
        return value

    try:
        from snaglist_pro.pipeline import correct_report_text

        corrected = correct_report_text(text)
        if corrected != text:
            return corrected
    except Exception:
        pass

    replacements = {
        "Emeregency": "Emergency",
        "allignment": "alignment",
        "temparary": "temporary",
        "Lighining": "Lighting",
        "achiving": "achieving",
        "Draiwings": "Drawings",
        "installedded": "installed",
        "installedd": "installed",
        "fixeded": "fixed",
        "obstuction": "obstruction",
        "scoket": "socket",
        "awarness": "awareness",
        "signgages": "signages",
        "engish": "English",
        "stips": "strips",
        "Exhuast": "Exhaust",
        "Diffueser": "Diffuser",
        "Landscaping": "Landscaping",
        "Lightining": "Lighting",
    }

    fixed = text
    for typo, corrected in replacements.items():
        fixed = fixed.replace(typo, corrected)
    return fixed


def _validate_export_workbook(workbook):
    if workbook is None:
        return ["Workbook is missing"]

    issues = []
    sheet_titles = {ws.title for ws in workbook.worksheets}
    if "Summary" not in sheet_titles:
        issues.append("Missing Summary sheet")

    data_sheets = [ws for ws in workbook.worksheets if ws.title not in {"Summary", "Correction Notes"}]
    if not data_sheets:
        issues.append("No data sheet found")

    for ws in data_sheets:
        if ws.max_row < 2:
            issues.append(f"Sheet '{ws.title}' is empty")
            break

    return issues


def _decode_archive_text(raw_bytes):
    if raw_bytes.startswith((b"\xff\xfe", b"\xfe\xff")):
        return raw_bytes.decode("utf-16")
    if raw_bytes.startswith(b"\xef\xbb\xbf"):
        return raw_bytes.decode("utf-8-sig")

    try:
        decoded = raw_bytes.decode("utf-8")
        if "\x00" not in decoded:
            return decoded
    except UnicodeDecodeError:
        pass

    for encoding in ("utf-16", "cp1252", "latin-1"):
        try:
            return raw_bytes.decode(encoding)
        except UnicodeDecodeError:
            continue
    return raw_bytes.decode("utf-8", errors="replace")


def _archive_text_candidates(archive):
    known_extensions = (".txt", ".csv", ".json", ".log", ".text", ".md")
    binary_extensions = (
        ".jpg", ".jpeg", ".png", ".gif", ".bmp", ".webp", ".mp4", ".avi", ".mov",
        ".pdf", ".doc", ".docx", ".xls", ".xlsx", ".dwg",
    )
    candidates = []

    for info in archive.infolist():
        name = info.filename
        lowered = name.lower()
        if info.is_dir() or lowered.endswith(binary_extensions):
            continue
        if lowered.endswith(known_extensions):
            candidates.append(name)
            continue
        if info.file_size > 5 * 1024 * 1024:
            continue
        try:
            sample = _decode_archive_text(archive.read(name)[:65536])
        except Exception:
            continue
        if parse_whatsapp_text(sample):
            candidates.append(name)

    return candidates


def _fallback_rows_from_text(text_content):
    rows = []
    for line in text_content.splitlines():
        description = line.strip()
        if description:
            rows.append({
                "Date": "",
                "Location": "",
                "Item": "WhatsApp note",
                "Description": description,
                "Category": "WhatsApp",
                "Status": "Open",
                "Priority": "Medium",
                "Source": "ZIP text",
                "Media Files": "",
                "Notes": "",
            })
    return rows


def _fallback_rows_from_media(media_names):
    return [
        {
            "Date": "",
            "Location": "",
            "Item": Path(name).name,
            "Description": "Image evidence from ZIP export",
            "Category": "Media",
            "Status": "Open",
            "Priority": "Medium",
            "Source": "ZIP media",
            "Media Files": Path(name).name,
            "Notes": "",
        }
        for name in media_names
    ]


def _report_image_names(media_names):
    return sorted(set(media_names))[:10]


def _prepare_image_for_excel(source_path, target_path, size_px, fit_inside):
    pixels = max(20, int(size_px))
    with Image.open(source_path) as image:
        image = image.convert("RGB")
        if fit_inside:
            image.thumbnail((pixels, pixels), Image.Resampling.LANCZOS)
        else:
            image = image.resize((pixels, pixels), Image.Resampling.LANCZOS)
        image.save(target_path, format="JPEG", quality=85, optimize=True)
    return target_path


def _normalize_image_size(value):
    if value is None:
        return 200
    text = str(value).strip().lower()
    if text.endswith("x"):
        text = text[:-1]
    if text in {"200x200", "200", "200px"}:
        return 200
    if text in {"150x150", "150", "150px"}:
        return 150
    if text in {"100x100", "100", "100px"}:
        return 100
    if text in {"80x80", "80", "80px"}:
        return 80
    return 200


def _fit_image_dimensions(width, height, max_size):
    width = max(1, int(width))
    height = max(1, int(height))
    max_size = max(20, int(max_size))
    scale = min(max_size / width, max_size / height, 1)
    return max(1, round(width * scale)), max(1, round(height * scale))


def _attach_image_to_cell(ws, row_index, column_letter, image_path, size_px=200, fit_inside=False):
    if not image_path:
        return False

    path = Path(image_path)
    if not path.exists():
        return False

    img = XLImage(str(path))
    pixels = max(20, int(size_px))
    if fit_inside:
        img.width, img.height = _fit_image_dimensions(img.width, img.height, pixels)
    else:
        img.width = pixels
        img.height = pixels

    ws.column_dimensions[column_letter].width = max(ws.column_dimensions[column_letter].width or 8, pixels / 7)
    ws.row_dimensions[row_index].height = max(ws.row_dimensions[row_index].height or 15, pixels * 0.75)
    ws.add_image(img, f"{column_letter}{row_index}")
    return True


def _first_value(mapping, candidate_keys, default=""):
    if not isinstance(mapping, dict):
        return default
    for key in candidate_keys:
        if key in mapping and mapping[key] not in (None, ""):
            return str(mapping[key]).strip()
        normalized_key = str(key).lower()
        for existing_key, value in mapping.items():
            if str(existing_key).strip().lower() == normalized_key and value not in (None, ""):
                return str(value).strip()
    return default


def _read_checklist_rows(file_path):
    path = Path(file_path)
    if not path.exists() or not path.is_file():
        return []

    suffix = path.suffix.lower()
    if suffix == ".csv":
        with path.open("r", encoding="utf-8-sig", newline="") as handle:
            import csv

            reader = csv.DictReader(handle)
            return [
                {str(k).strip(): (v if v is not None else "") for k, v in row.items()}
                for row in reader
                if row
            ]

    if suffix in {".xlsx", ".xlsm"}:
        try:
            workbook = load_workbook(path, read_only=True, data_only=True)
            if not workbook.worksheets:
                return []
            sheet = workbook.worksheets[0]
            rows = list(sheet.iter_rows(values_only=True))
            if not rows:
                return []
            header_index = None
            for index, row in enumerate(rows[:20]):
                normalized = {str(cell).strip().lower() for cell in row if cell is not None and str(cell).strip()}
                if {"slno", "check points"}.issubset(normalized) or "check points" in normalized:
                    header_index = index
                    break
            if header_index is None:
                header_index = 0

            headers = [str(cell).strip() if cell is not None else "" for cell in rows[header_index]]
            parsed = []
            for row in rows[header_index + 1:]:
                if not any(cell is not None and str(cell).strip() for cell in row):
                    continue
                item = {}
                for index, header in enumerate(headers):
                    if header and index < len(row):
                        item[header] = row[index]
                parsed.append(item)
            return parsed
        except Exception:
            return []

    return []


def _build_export_rows(chat_rows, checklist_rows):
    checklist_rows = checklist_rows or []
    if checklist_rows:
        export_rows = []
        for item in checklist_rows:
            if not isinstance(item, dict):
                continue
            row = {
                "Date": _first_value(item, ["Date", "Created Date", "Inspection Date", "date"], ""),
                "Location": _first_value(item, ["Location", "Location Name", "Area", "Area/location", "Room", "Zone", "room", "location"], ""),
                "Item": _first_value(item, ["Item", "Title", "Component", "Name", "Check", "Check points", "Check Points", "item"], ""),
                "Description": _first_value(item, ["Description", "Issue", "Defect", "Comment", "Details", "Snagpoints", "Snag Points", "description", "Check points"], ""),
                "Category": _first_value(item, ["Category", "Catagory", "Trade", "Section", "Type", "category"], "General"),
                "Status": _first_value(item, ["Status", "Status (Open/Close)", "State", "Result", "status"], "Open"),
                "Priority": _first_value(item, ["Priority", "Severity", "Urgency", "priority"], "Medium"),
                "Source": "Checklist",
                "Media Files": _first_value(item, ["Media Files", "Media", "Attachments", "media_files"], ""),
                "Notes": _first_value(item, ["Notes", "Remarks", "Observation", "notes"], ""),
            }
            if not row["Description"]:
                row["Description"] = _first_value(item, ["Issue Description", "Description of Defect", "Defect Details", "issue"], "")
            if not row["Item"] and row["Description"]:
                row["Item"] = row["Description"]
            export_rows.append(row)
        return export_rows

    export_rows = []
    for item in chat_rows or []:
        if not isinstance(item, dict):
            continue
        clean_message = str(item.get("message") or "").strip()
        row = {
            "Date": _first_value(item, ["date", "Date"], ""),
            "Location": _first_value(item, ["location", "Location", "room", "Room", "area", "Area"], ""),
            "Item": _first_value(item, ["item", "Item", "title", "Title"], "WhatsApp note"),
            "Description": clean_message,
            "Category": "WhatsApp",
            "Status": "Open",
            "Priority": "Medium",
            "Source": "WhatsApp",
            "Media Files": "; ".join(str(v) for v in item.get("media_files", []) if v),
            "Notes": item.get("sender") or "",
        }
        if not row["Date"] and item.get("time"):
            row["Date"] = item.get("time")
        export_rows.append(row)
    return export_rows


class SnaglistDesktopApp:
    def __init__(self, root):
        self.root = root
        self.root.title(WINDOW_TITLE)
        self.root.geometry("920x620")
        self.root.minsize(820, 540)
        self.root.configure(bg="#0f172a")
        self.root.option_add("*Font", "{Segoe UI} 10")

        self.current_user = None
        self.selected_zip = None
        self.selected_checklist = None
        self.latest_output_file = None
        self.latest_output_bytes = None
        self.latest_suggested_name = None
        self.download_btn = None
        self.status_var = tk.StringVar(value="Idle")
        self.progress_var = tk.DoubleVar(value=0)
        self.stat_vars = {}
        self.image_size_var = tk.StringVar(value="200x200")
        self.fit_inside_cell_var = tk.BooleanVar(value=True)
        self.spelling_correction_var = tk.BooleanVar(value=True)
        self.root.protocol("WM_DELETE_WINDOW", self._on_close)
        _setup_file_logging()
        self._apply_modern_theme()
        self._build_login_screen()
        self._load_session()

    def _apply_modern_theme(self):
        self._apply_theme_colors("#0f172a", "#f8fafc")

    def _apply_theme_colors(self, dark_bg, light_bg):
        style = ttk.Style(self.root)
        style.theme_use("clam")

        style.configure("Modern.TFrame", background=light_bg)
        style.configure("Modern.TLabel", background=light_bg, foreground=dark_bg)
        style.configure("ModernAccent.TLabel", background=dark_bg, foreground=light_bg)
        style.configure("Modern.TButton", padding=(16, 10), relief="flat", background="#2563eb", foreground="#ffffff", font=_safe_font(10, "bold"))
        style.map("Modern.TButton", background=[("active", "#1d4ed8")])
        style.configure("Secondary.TButton", padding=(16, 10), relief="flat", background="#e2e8f0", foreground=dark_bg, font=_safe_font(10, "bold"))
        style.map("Secondary.TButton", background=[("active", "#cbd5e1")])
        style.configure("Modern.TEntry", fieldbackground="#ffffff", foreground=dark_bg, bordercolor="#cbd5e1")
        style.configure("Header.TLabel", background=dark_bg, foreground=light_bg, font=_safe_font(26, "bold"))
        style.configure("SubHeader.TLabel", background=dark_bg, foreground="#94a3b8", font=_safe_font(11))
        style.configure("Card.TFrame", background="#ffffff")
        style.configure("Card.TLabel", background="#ffffff", foreground="#334155")
        style.configure("Modern.TCheckbutton", background=light_bg)
        self._dark_bg = dark_bg
        self._light_bg = light_bg

    def _toggle_theme(self):
        if getattr(self, "_is_dark", True):
            self._apply_theme_colors("#f8fafc", "#0f172a")
            self._is_dark = False
            self.root.configure(bg="#f8fafc")
        else:
            self._apply_theme_colors("#0f172a", "#f8fafc")
            self._is_dark = True
            self.root.configure(bg="#0f172a")
        self._refresh_theme()

    def _refresh_theme(self):
        if not getattr(self, "_is_dark", True):
            for widget in self.root.winfo_children():
                self._refresh_widget_theme(widget)

    def _refresh_widget_theme(self, widget):
        try:
            bg = widget.cget("bg")
            fg = widget.cget("fg")
            font = widget.cget("font")
            if bg == "#0f172a":
                widget.configure(bg="#f8fafc")
            elif bg == "#f8fafc":
                widget.configure(bg="#0f172a")
            if fg == "#f8fafc":
                widget.configure(fg="#0f172a")
            elif fg == "#0f172a":
                widget.configure(fg="#f8fafc")
        except Exception:
            pass
        for child in widget.winfo_children():
            self._refresh_widget_theme(child)

    def _bind_keyboard_shortcuts(self):
        self.root.bind_all("<Control-o>", lambda e: self.choose_zip())
        self.root.bind_all("<Control-g>", lambda e: self.choose_checklist())
        self.root.bind_all("<Control-s>", lambda e: self.download_output())
        self.root.bind_all("<Control-t>", lambda e: self._toggle_theme())

    def _update_summary_formulas(self, workbook, total):
        summary = workbook["Summary"]
        if total:
            summary["B8"] = "=B6/B5"
        else:
            summary["B8"] = 0
        if total:
            summary["B8"].number_format = "0%"

    def _save_session(self):
        try:
            app_dir = Path(os.environ.get("LOCALAPPDATA", Path.home())) / "SnaglistPro"
            app_dir.mkdir(parents=True, exist_ok=True)
            session = {
                "selected_zip": self.selected_zip if getattr(self, "selected_zip", None) else "",
                "selected_checklist": self.selected_checklist if getattr(self, "selected_checklist", None) else "",
            }
            (app_dir / "session.json").write_text(json.dumps(session), encoding="utf-8")
        except Exception:
            pass

    def _load_session(self):
        try:
            app_dir = Path(os.environ.get("LOCALAPPDATA", Path.home())) / "SnaglistPro"
            session_file = app_dir / "session.json"
            if session_file.exists():
                session = json.loads(session_file.read_text(encoding="utf-8"))
                zip_path = session.get("selected_zip", "")
                checklist_path = session.get("selected_checklist", "")
                if zip_path and Path(zip_path).exists():
                    self.selected_zip = zip_path
                if checklist_path and Path(checklist_path).exists():
                    self.selected_checklist = checklist_path
        except Exception:
            pass

    def _build_batch_tab(self):
        batch_win = tk.Toplevel(self.root)
        batch_win.title("Batch Processing")
        batch_win.geometry("640x480")
        batch_win.resizable(True, True)

        tk.Label(batch_win, text="Batch ZIP Processing", font=_safe_font(16, "bold"), bg="#f8fafc", fg="#0f172a").pack(pady=12)

        tk.Label(batch_win, text="Select multiple WhatsApp ZIP exports:", bg="#f8fafc", fg="#475569", font=_safe_font(10)).pack()

        self.batch_zip_listbox = tk.Listbox(batch_win, height=8, width=60, font=_safe_font(10))
        self.batch_zip_listbox.pack(pady=8, padx=20)

        def add_batch_zips():
            paths = filedialog.askopenfilenames(
                title="Select WhatsApp ZIP exports",
                filetypes=[("ZIP Files", "*.zip")],
            )
            for p in paths:
                self.batch_zip_listbox.insert(tk.END, p)

        def remove_batch_zip():
            sel = self.batch_zip_listbox.curselection()
            if sel:
                self.batch_zip_listbox.delete(sel[0])

        btn_frame = tk.Frame(batch_win, bg="#f8fafc")
        btn_frame.pack(pady=6)
        tk.Button(btn_frame, text="Add ZIPs", command=add_batch_zips, bg="#2563eb", fg="#ffffff", font=_safe_font(10, "bold"), padx=12, pady=6).pack(side="left", padx=6)
        tk.Button(btn_frame, text="Remove", command=remove_batch_zip, bg="#e2e8f0", fg="#0f172a", font=_safe_font(10, "bold"), padx=12, pady=6).pack(side="left", padx=6)

        self.batch_status = tk.StringVar(value="Ready")
        tk.Label(batch_win, textvariable=self.batch_status, bg="#f8fafc", fg="#2563eb", font=_safe_font(10, "bold")).pack(pady=6)

        def run_batch():
            zip_paths = self.batch_zip_listbox.get(0, tk.END)
            if not zip_paths:
                messagebox.showwarning("No ZIPs", "Add at least one ZIP file to process.")
                return
            self.batch_status.set(f"Processing {len(zip_paths)} files...")
            try:
                from snaglist_pro.batch_processor import process_batch
                checklist = self.selected_checklist if getattr(self, "selected_checklist", None) else None
                output_dir = str(Path(os.environ.get("LOCALAPPDATA", Path.home())) / "SnaglistPro" / "batch_output")
                results = process_batch(list(zip_paths), checklist_path=checklist, output_base=output_dir)
                success = sum(1 for r in results if r.get("success"))
                failed = sum(1 for r in results if not r.get("success"))
                self.batch_status.set(f"Done. {success} succeeded, {failed} failed.")
                messagebox.showinfo("Batch Complete", f"Processing complete:\n{success} succeeded, {failed} failed.\nOutput: {output_dir}")
            except Exception as exc:
                self.batch_status.set("Failed")
                messagebox.showerror("Batch failed", str(exc))

        tk.Button(batch_win, text="Run Batch", command=run_batch, bg="#10b981", fg="#ffffff", font=_safe_font(10, "bold"), padx=18, pady=10).pack(pady=10)

    def _build_login_screen(self):
        self.clear_widgets()

        shell = tk.Frame(self.root, bg="#0f172a")
        shell.pack(fill="both", expand=True)

        left = tk.Frame(shell, bg="#0f172a", width=440)
        left.pack(side="left", fill="y", padx=(0, 0))
        left.pack_propagate(False)

        branding = tk.Frame(left, bg="#0f172a")
        branding.pack(anchor="w", padx=48, pady=(60, 20))

        title = tk.Label(branding, text="Snaglist Pro", bg="#0f172a", fg="#f8fafc", font=_safe_font(30, "bold"))
        title.pack(anchor="w")
        sub = tk.Label(branding, text="Modern defect and snag tracking", bg="#0f172a", fg="#94a3b8", font=_safe_font(12))
        sub.pack(anchor="w", pady=(8, 0))

        right = tk.Frame(shell, bg="#f8fafc")
        right.pack(side="right", fill="both", expand=True, padx=40)

        form = tk.Frame(right, bg="#f8fafc")
        form.pack(expand=True, padx=60, pady=40)

        tk.Label(form, text="Login", bg="#f8fafc", fg="#0f172a", font=_safe_font(26, "bold")).pack(anchor="w", pady=(0, 18))
        tk.Label(form, text="Username", bg="#f8fafc", fg="#475569", font=_safe_font(10, "bold")).pack(anchor="w")
        self.username_var = tk.StringVar(value=DEFAULT_USERNAME)
        username_entry = tk.Entry(form, textvariable=self.username_var, width=32, font=_safe_font(11), bd=1, relief="solid", highlightthickness=1, highlightbackground="#cbd5e1", highlightcolor="#2563eb")
        username_entry.pack(fill="x", pady=(6, 12))

        tk.Label(form, text="Password", bg="#f8fafc", fg="#475569", font=_safe_font(10, "bold")).pack(anchor="w")
        self.password_var = tk.StringVar(value=DEFAULT_PASSWORD)
        password_entry = tk.Entry(form, textvariable=self.password_var, width=32, show="*", font=_safe_font(11), bd=1, relief="solid", highlightthickness=1, highlightbackground="#cbd5e1", highlightcolor="#2563eb")
        password_entry.pack(fill="x", pady=(6, 18))

        login_btn = tk.Button(form, text="Login", command=self.handle_login, bg="#2563eb", fg="#ffffff", activebackground="#1d4ed8", font=_safe_font(10, "bold"), bd=0, padx=18, pady=10, width=20)
        login_btn.pack(fill="x", pady=(6, 0))

        watermark = tk.Label(
            self.root,
            text="Developed by Vineeth.Bisa",
            fg="#7e7e7e",
            bg="#0f172a",
            font=_safe_font(9),
        )
        watermark.pack(side="bottom", anchor="se", padx=12, pady=8)
        self._bind_keyboard_shortcuts()
        self._is_dark = True

    def _build_main_screen(self):
        self.clear_widgets()

        shell = tk.Frame(self.root, bg="#eef2ff")
        shell.pack(fill="both", expand=True)

        top = tk.Frame(shell, bg="#0f172a", height=130)
        top.pack(fill="x")
        top.pack_propagate(False)

        title = tk.Label(top, text="Snaglist Pro", bg="#0f172a", fg="#f8fafc", font=_safe_font(27, "bold"))
        title.place(x=26, y=24)

        sub = tk.Label(top, text="Construction defect and snag workflow", bg="#0f172a", fg="#94a3b8", font=_safe_font(11))
        sub.place(x=26, y=72)

        actions = tk.Frame(shell, bg="#eef2ff")
        actions.pack(fill="x", padx=20, pady=(18, 12))

        choose_zip_btn = tk.Button(actions, text="Choose ZIP", command=self.choose_zip, bg="#2563eb", fg="#ffffff", activebackground="#1d4ed8", font=_safe_font(10, "bold"), bd=0, padx=18, pady=10)
        choose_zip_btn.pack(side="left")

        choose_checklist_btn = tk.Button(actions, text="Choose Checklist", command=self.choose_checklist, bg="#0f172a", fg="#ffffff", activebackground="#111827", font=_safe_font(10, "bold"), bd=0, padx=18, pady=10)
        choose_checklist_btn.pack(side="left", padx=(10, 0))

        size_label = tk.Label(actions, text="Image size", bg="#eef2ff", fg="#334155", font=_safe_font(10, "bold"))
        size_label.pack(side="left", padx=(10, 6))
        self.image_size_combo = ttk.Combobox(actions, textvariable=self.image_size_var, values=IMAGE_SIZE_OPTIONS, state="readonly", width=12)
        self.image_size_combo.pack(side="left")

        self.fit_inside_cell_check = ttk.Checkbutton(actions, text=FIT_INSIDE_CELL_LABEL, variable=self.fit_inside_cell_var)
        self.fit_inside_cell_check.pack(side="left", padx=(10, 0))

        self.spelling_check = ttk.Checkbutton(actions, text=SPELLING_CORRECTION_LABEL, variable=self.spelling_correction_var)
        self.spelling_check.pack(side="left", padx=(10, 0))

        process_btn = tk.Button(actions, text="Generate Snaglist", command=self.process_zip, bg="#0f172a", fg="#ffffff", activebackground="#111827", font=_safe_font(10, "bold"), bd=0, padx=18, pady=10)
        process_btn.pack(side="left", padx=(10, 0))

        self.download_btn = tk.Button(
            actions,
            text="Download Excel/Sheets",
            command=self.download_output,
            bg="#10b981",
            fg="#ffffff",
            activebackground="#059669",
            font=_safe_font(10, "bold"),
            bd=0,
            padx=18,
            pady=10,
            state="disabled",
        )
        self.download_btn.pack(side="left", padx=(10, 0))

        logout_btn = tk.Button(actions, text="Logout", command=self.logout, bg="#e2e8f0", fg="#0f172a", activebackground="#cbd5e1", font=_safe_font(10, "bold"), bd=0, padx=18, pady=10)
        logout_btn.pack(side="right")

        batch_btn = tk.Button(actions, text="Batch Process", command=self._build_batch_tab, bg="#7c3aed", fg="#ffffff", activebackground="#6d28d9", font=_safe_font(10, "bold"), bd=0, padx=18, pady=10)
        batch_btn.pack(side="right", padx=(10, 0))

        theme_btn = tk.Button(actions, text="Theme", command=self._toggle_theme, bg="#475569", fg="#ffffff", activebackground="#334155", font=_safe_font(10, "bold"), bd=0, padx=14, pady=10)
        theme_btn.pack(side="right", padx=(10, 0))

        status_frame = tk.Frame(shell, bg="#ffffff", highlightthickness=1, highlightbackground="#dfe7f5")
        status_frame.pack(fill="x", padx=20, pady=(0, 12))
        tk.Label(status_frame, text="Progress status", bg="#ffffff", fg="#0f172a", font=_safe_font(10, "bold"), anchor="w").pack(anchor="w", padx=14, pady=(12, 4))
        self.status_var.set("Ready")
        self.progress_var.set(0)
        status_label = tk.Label(status_frame, textvariable=self.status_var, bg="#ffffff", fg="#2563eb", font=_safe_font(11, "bold"), anchor="w")
        status_label.pack(anchor="w", padx=14, pady=(0, 10))
        progress = ttk.Progressbar(status_frame, variable=self.progress_var, maximum=100)
        progress.pack(fill="x", padx=14, pady=(0, 12))

        stats = tk.Frame(shell, bg="#eef2ff")
        stats.pack(fill="x", padx=20, pady=(0, 12))

        self.stat_vars = {}
        for idx, (label_text, value_text, accent) in enumerate([
            ("Open Snags", "0", "#2563eb"),
            ("In Progress", "0", "#f59e0b"),
            ("Resolved", "0", "#10b981"),
        ]):
            card = tk.Frame(stats, bg="#ffffff", highlightthickness=1, highlightbackground="#dfe7f5", width=200, height=94)
            card.grid(row=0, column=idx, padx=(0 if idx == 0 else 12, 0), sticky="nsew")
            card.grid_propagate(False)

            accent_bar = tk.Frame(card, bg=accent, height=6)
            accent_bar.pack(fill="x")

            tk.Label(card, text=label_text, bg="#ffffff", fg="#64748b", font=_safe_font(10), anchor="w").pack(anchor="w", padx=14, pady=(12, 4))
            value_label = tk.Label(card, text=value_text, bg="#ffffff", fg="#0f172a", font=_safe_font(22, "bold"), anchor="w")
            value_label.pack(anchor="w", padx=14)
            self.stat_vars[label_text] = value_label

        stats.grid_columnconfigure((0, 1, 2), weight=1)

        card = tk.Frame(shell, bg="#ffffff", highlightthickness=1, highlightbackground="#e2e8f0")
        card.pack(fill="x", padx=20, pady=(0, 14))

        self.zip_label = tk.Label(card, text="No ZIP selected", anchor="w", justify="left", bg="#ffffff", fg="#334155", font=_safe_font(10), padx=16, pady=14)
        self.zip_label.pack(fill="x")

        checklist_card = tk.Frame(shell, bg="#ffffff", highlightthickness=1, highlightbackground="#e2e8f0")
        checklist_card.pack(fill="x", padx=20, pady=(0, 14))
        self.checklist_label = tk.Label(checklist_card, text="No checklist selected", anchor="w", justify="left", bg="#ffffff", fg="#334155", font=_safe_font(10), padx=16, pady=14)
        self.checklist_label.pack(fill="x")

        log_frame = tk.Frame(shell, bg="#ffffff", highlightthickness=1, highlightbackground="#e2e8f0")
        log_frame.pack(fill="both", expand=True, padx=20, pady=(0, 18))

        title_word = tk.Label(log_frame, text="Processing log", bg="#ffffff", fg="#0f172a", font=_safe_font(11, "bold"), anchor="w")
        title_word.pack(anchor="w", padx=14, pady=(12, 6))

        self.log_box = tk.Text(log_frame, height=14, wrap="word", bd=0, padx=12, pady=12, bg="#f8fafc", fg="#0f172a", font=_safe_font(10))
        self.log_box.pack(fill="both", expand=True, padx=12, pady=(0, 12))
        self.log_box.insert("end", "Ready. Select a WhatsApp ZIP export and checklist to begin.\n")
        self.log_box.config(state="disabled")

        watermark = tk.Label(
            self.root,
            text="Developed by Vineeth.Bisa",
            fg="#7e7e7e",
            bg="#eef2ff",
            font=_safe_font(9),
        )
        watermark.pack(side="bottom", anchor="se", padx=12, pady=8)

    def clear_widgets(self):
        for widget in self.root.winfo_children():
            widget.destroy()

    def handle_login(self):
        username = self.username_var.get().strip()
        password = self.password_var.get()

        from snaglist_pro.auth import verify_credentials, get_configured_user
        if get_configured_user() is not None:
            if verify_credentials(username, password):
                self.current_user = username
                self._build_main_screen()
                return
            messagebox.showerror("Login failed", "Invalid username or password.")
            return

        if username == DEFAULT_USERNAME and password == DEFAULT_PASSWORD:
            self.current_user = username
            self._build_main_screen()
            return

        messagebox.showerror("Login failed", "Invalid username or password.")

    def logout(self):
        self._save_session()
        self.current_user = None
        self._build_login_screen()

    def _cleanup_temp_files(self):
        try:
            export_dir = _get_output_dir()
            if export_dir.exists():
                import shutil as _shutil
                _shutil.rmtree(export_dir, ignore_errors=True)
        except Exception:
            pass

    def _on_close(self):
        self._save_session()
        self._cleanup_temp_files()
        self.root.destroy()

    def choose_zip(self):
        file_path = filedialog.askopenfilename(
            title="Select WhatsApp ZIP export",
            filetypes=[("ZIP Files", "*.zip")],
        )
        if file_path:
            self.selected_zip = file_path
            self.zip_label.config(text=file_path)
            self.append_log(f"Selected ZIP: {file_path}")

    def choose_checklist(self):
        file_path = filedialog.askopenfilename(
            title="Select checklist file",
            filetypes=[("Excel Files", "*.xlsx"), ("CSV Files", "*.csv"), ("All files", "*.*")],
        )
        if file_path:
            self.selected_checklist = file_path
            self.checklist_label.config(text=file_path)
            self.append_log(f"Selected checklist: {file_path}")

    def _set_dashboard_counts(self, total=0, open_count=0, in_progress=0, resolved=0):
        self.stat_vars.get("Open Snags", self.stat_vars.get("Open Snags")).config(text=str(open_count))
        self.stat_vars.get("In Progress", self.stat_vars.get("In Progress")).config(text=str(in_progress))
        self.stat_vars.get("Resolved", self.stat_vars.get("Resolved")).config(text=str(resolved))

    def _update_progress(self, status, pct):
        self.status_var.set(status)
        self.progress_var.set(pct)
        self.root.update_idletasks()

    def process_zip(self):
        if not getattr(self, "selected_zip", None):
            messagebox.showwarning("No ZIP selected", "Please choose a ZIP file first.")
            return

        self.latest_output_file = None
        self.latest_output_bytes = None
        self.latest_suggested_name = None
        if self.download_btn is not None:
            self.download_btn.config(state="disabled")

        self._update_progress("Preparing ZIP", 10)
        self.append_log("Processing ZIP file...")
        self.append_log(f"Checklist file: {self.selected_checklist}")
        self.root.update_idletasks()

        file_bytes, suggested_name = self._generate_output_from_zip(self.selected_zip)
        if file_bytes:
            self.latest_output_file = suggested_name
            self.latest_output_bytes = file_bytes
            self.latest_suggested_name = suggested_name
            if self.download_btn is not None:
                self.download_btn.config(state="normal")
            self._update_progress("Completed", 100)
            self._set_dashboard_counts(open_count=1, in_progress=1, resolved=0)
            self.append_log(f"Output generated ({len(file_bytes):,} bytes)")
            messagebox.showinfo("Output ready", f"Snaglist ready for download:\n{suggested_name}")
        else:
            self._update_progress("Failed", 0)
            self.append_log("No readable WhatsApp content found in the ZIP.")
            messagebox.showerror("Processing failed", "No readable WhatsApp data was found in the selected ZIP file.")

    def _generate_output_from_zip(self, zip_path):
        try:
            from io import BytesIO
            from snaglist_pro.pipeline import SnaglistPipeline

            output_dir = _get_output_dir()
            output_dir.mkdir(parents=True, exist_ok=True)
            selected_checklist = getattr(self, "selected_checklist", None)
            if selected_checklist and Path(selected_checklist).suffix.lower() not in {".csv", ".xlsx", ".xlsm"}:
                selected_checklist = None

            def pipeline_progress(_step, message, percent):
                self._update_progress(message, percent)

            try:
                result = SnaglistPipeline().run(
                    zip_path,
                    selected_checklist,
                    str(output_dir),
                    fast_mode=False,
                    ai_enabled=False,
                    progress_callback=pipeline_progress,
                )
                output_file = result.get("excel_path") if isinstance(result, dict) else None
                if output_file:
                    self.append_log(f"Canonical pipeline output: {output_file}")
                    path = Path(output_file)
                    file_bytes = path.read_bytes()
                    suggested_name = path.name
                    return file_bytes, suggested_name
            except Exception as exc:
                self.append_log(f"Canonical pipeline fallback: {exc}")
                with zipfile.ZipFile(zip_path, "r") as archive:
                    chat_entries = [
                        name for name in archive.namelist()
                        if name.lower().endswith((".txt", ".csv", ".json", ".log"))
                    ]
                    chat_rows = []
                    for name in chat_entries:
                        chat_rows.extend(parse_whatsapp_text(_decode_archive_text(archive.read(name))))
                if chat_rows:
                    raise RuntimeError(
                        "The canonical snaglist pipeline failed for a readable WhatsApp export; "
                        "the legacy exporter was blocked to prevent an incorrect workbook."
                    ) from exc

            self._update_progress("Reading ZIP content", 25)
            with zipfile.ZipFile(zip_path, "r") as archive:
                text_files = _archive_text_candidates(archive)
                media_names = [
                    info.filename for info in archive.infolist()
                    if info.filename.lower().endswith((".jpg", ".jpeg", ".png", ".gif", ".bmp", ".webp"))
                ]
                self.append_log(f"ZIP entries: {len(archive.infolist())}; text candidates: {len(text_files)}")

                checklist_rows = []
                if selected_checklist:
                    checklist_rows = _read_checklist_rows(selected_checklist)

                if not text_files and not checklist_rows and not media_names:
                    return None, None

                all_rows = []
                raw_texts = []
                total = len(text_files)
                for index, name in enumerate(text_files, start=1):
                    self._update_progress(f"Parsing file {index}/{total}", 25 + int((index / total) * 40))
                    with archive.open(name) as handle:
                        raw = _decode_archive_text(handle.read())
                    raw_texts.append(raw)
                    parsed_rows = parse_whatsapp_text(raw)
                    all_rows.extend(parsed_rows)
                self.append_log(f"Parsed WhatsApp messages: {len(all_rows)}")

                if not all_rows and not checklist_rows:
                    return None, None

                export_rows = _build_export_rows(all_rows, checklist_rows)
                if not export_rows:
                    for raw in raw_texts:
                        export_rows.extend(_fallback_rows_from_text(raw))
                if not export_rows:
                    export_rows = _fallback_rows_from_media(media_names)
                if not export_rows:
                    return None, None

                self._update_progress("Writing report", 75)
                timestamp = datetime.now().strftime("%Y%m%d_%H%M%S")
                suggested_name = f"snaglist_report_{timestamp}.xlsx"

                workbook = Workbook()
                sheet = workbook.active
                sheet.title = "Snaglist"

                spelling_toggle = getattr(self, "spelling_correction_var", None)
                spelling_enabled = bool(spelling_toggle.get() if spelling_toggle is not None else True)
                for row in export_rows:
                    for field in ("Description", "Item", "Location", "Notes"):
                        if field in row:
                            row[field] = _apply_optional_spelling_correction(row.get(field), spelling_enabled)

                sheet.append(EXPORT_COLUMNS)
                for row in export_rows:
                    sheet.append([row.get(column, "") for column in EXPORT_COLUMNS])

                summary = workbook.create_sheet("Summary")
                summary.append(["Metric", "Value"])
                summary.append(["Total Items", len(export_rows)])
                summary.append(["Open", sum(1 for row in export_rows if str(row.get("Status", "")).strip().lower() in {"open", "new"})])
                summary.append(["In Progress", sum(1 for row in export_rows if str(row.get("Status", "")).strip().lower() in {"in progress", "progress", "pending"})])
                summary.append(["Resolved", sum(1 for row in export_rows if str(row.get("Status", "")).strip().lower() in {"closed", "resolved", "done"})])

                if spelling_enabled:
                    notes = workbook.create_sheet("Correction Notes")
                    notes.append(["Original", "Corrected", "Field"])

                validation_issues = _validate_export_workbook(workbook)
                if validation_issues:
                    message = "; ".join(validation_issues)
                    self.append_log(f"Workbook validation failed: {message}")
                    raise ValueError(message)

                image_size_var = getattr(self, "image_size_var", None)
                fit_inside_var = getattr(self, "fit_inside_cell_var", None)
                image_size = _normalize_image_size(image_size_var.get() if image_size_var is not None else "200x200")
                fit_inside_cell = bool(fit_inside_var.get() if fit_inside_var is not None else True)
                image_dir = output_dir / "temp_images"
                image_dir.mkdir(parents=True, exist_ok=True)

                report_images = _report_image_names(
                    item for item in archive.namelist()
                    if item.lower().endswith((".jpg", ".jpeg", ".png", ".gif", ".bmp", ".webp"))
                )
                for idx, name in enumerate(report_images):
                    source_target = image_dir / Path(name).name
                    target = image_dir / f"report_image_{idx + 1}.jpg"
                    try:
                        with archive.open(name) as src, target.open("wb") as dst:
                            dst.write(src.read())
                        _prepare_image_for_excel(str(target), str(source_target), image_size, fit_inside_cell)
                    except Exception:
                        continue
                    _attach_image_to_cell(sheet, idx + 2, "K", str(source_target), image_size, fit_inside_cell)

                for column_cells in sheet.columns:
                    max_length = max(len(str(cell.value)) if cell.value is not None else 0 for cell in column_cells)
                    sheet.column_dimensions[column_cells[0].column_letter].width = min(max_length + 2, 30)

                buf = BytesIO()
                workbook.save(buf)
                buf.seek(0)
                file_bytes = buf.read()
                return file_bytes, suggested_name
        except Exception as exc:
            self._update_progress("Error", 0)
            messagebox.showerror("ZIP processing failed", str(exc))
            return None, None

    def download_output(self):
        file_bytes = getattr(self, "latest_output_bytes", None)
        if not file_bytes:
            messagebox.showinfo("No file ready", "Generate a snaglist first before downloading.")
            return

        suggested = self.latest_suggested_name or "snaglist_report.xlsx"

        target_path = filedialog.asksaveasfilename(
            title="Save snaglist output",
            defaultextension=".xlsx",
            initialfile=_suggest_download_filename(suggested),
            filetypes=[("Excel/Google Sheets Files", "*.xlsx"), ("CSV Files", "*.csv"), ("All files", "*.*")],
        )
        if not target_path:
            return

        try:
            with open(target_path, "wb") as dst:
                dst.write(file_bytes)
            self.append_log(f"Downloaded output to: {target_path}")
            messagebox.showinfo("Download complete", f"File saved to:\n{target_path}")
        except Exception as exc:
            messagebox.showerror("Download failed", str(exc))

    def append_log(self, message):
        if not hasattr(self, "log_box"):
            return
        self.log_box.config(state="normal")
        self.log_box.insert("end", message + "\n")
        self.log_box.see("end")
        self.log_box.config(state="disabled")


def run_desktop_app():
    root = tk.Tk()
    app = SnaglistDesktopApp(root)
    root.mainloop()


if __name__ == "__main__":
    run_desktop_app()
