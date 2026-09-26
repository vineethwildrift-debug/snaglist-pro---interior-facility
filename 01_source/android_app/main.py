from __future__ import annotations

import os
import zipfile
from pathlib import Path

from openpyxl import Workbook, load_workbook

from desktop_app import _apply_optional_spelling_correction, _archive_text_candidates, _build_export_rows, _decode_archive_text, _read_checklist_rows, _validate_export_workbook
from snaglist_pro.parse_export import parse_whatsapp_text


def create_android_app():
    return {"name": "Snaglist Pro", "platform": "android", "status": "ready"}


def generate_report_from_uploads(zip_path: str, checklist_path: str, correct_spelling: bool = True):
    if not zip_path or not checklist_path:
        raise ValueError("ZIP export and checklist are required.")

    zip_file = Path(zip_path)
    checklist_file = Path(checklist_path)

    if not zip_file.exists():
        raise ValueError("ZIP export was not found.")
    if not checklist_file.exists():
        raise ValueError("Checklist file was not found.")

    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Snaglist"
    rows = []

    with zipfile.ZipFile(zip_file, "r") as archive:
        text_files = _archive_text_candidates(archive)
        parsed_rows = []
        for name in text_files:
            try:
                raw_text = _decode_archive_text(archive.read(name))
                parsed_rows.extend(parse_whatsapp_text(raw_text))
            except Exception:
                continue
            
        checklist_rows = _read_checklist_rows(str(checklist_file))
        rows = _build_export_rows(parsed_rows, checklist_rows)

    if not rows:
        raise ValueError("No valid data rows were found in the ZIP or checklist.")

    sheet.append([
        "Date", "Location", "Item", "Description", "Category", "Status", "Priority",
        "Source", "Media Files", "Notes",
    ])

    for row in rows:
        normalized = {}
        for key in [
            "Date", "Location", "Item", "Description", "Category", "Status", "Priority",
            "Source", "Media Files", "Notes",
        ]:
            value = row.get(key, "")
            if correct_spelling and key in {"Description", "Item", "Location", "Notes"}:
                value = _apply_optional_spelling_correction(value, True)
            normalized[key] = value

        sheet.append([normalized.get(column, "") for column in [
            "Date", "Location", "Item", "Description", "Category", "Status", "Priority",
            "Source", "Media Files", "Notes",
        ]])

    summary = workbook.create_sheet("Summary")
    summary.append(["Metric", "Value"])
    summary.append(["Total Items", len(rows)])
    summary.append(["Open", sum(1 for row in rows if str(row.get("Status", "")).strip().lower() in {"open", "new"})])
    summary.append(["In Progress", sum(1 for row in rows if str(row.get("Status", "")).strip().lower() in {"in progress", "pending"})])
    summary.append(["Resolved", sum(1 for row in rows if str(row.get("Status", "")).strip().lower() in {"closed", "resolved", "done"})])

    issues = _validate_export_workbook(workbook)
    if issues:
        raise ValueError("Workbook validation failed: " + "; ".join(issues))

    output_dir = Path(os.environ.get("LOCALAPPDATA", str(Path.home()))) / "SnaglistProAndroid" / "exports"
    output_dir.mkdir(parents=True, exist_ok=True)
    output_path = output_dir / f"snaglist_android_{abs(hash(str(zip_path) + str(checklist_path)))}.xlsx"
    workbook.save(output_path)
    return str(output_path)
